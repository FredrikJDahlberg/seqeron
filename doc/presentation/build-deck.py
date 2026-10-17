"""Builds the seqeron architecture briefing deck."""
import sys
import uuid
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.dml import MSO_THEME_COLOR as T, MSO_LINE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

OUT = sys.argv[1]
import os
OFF = set(os.environ.get("OFF", "").split(","))

TX, BG, DK2, LT2 = T.TEXT_1, T.BACKGROUND_1, T.TEXT_2, T.BACKGROUND_2
A1, A2, A3, A4, A5, A6 = T.ACCENT_1, T.ACCENT_2, T.ACCENT_3, T.ACCENT_4, T.ACCENT_5, T.ACCENT_6
# A1 teal (sequenced, live)  A2 amber (synthesized, history)  A3 teal tint  A4 amber tint  A5 muted  A6 rule
COLORS = {"dk1": "15202B", "lt1": "FFFFFF", "dk2": "0F2A37", "lt2": "EEF2F5",
          "accent1": "0D5C78", "accent2": "B0601C", "accent3": "D9EAF0", "accent4": "F6E5D4",
          "accent5": "56636F", "accent6": "C9D2DA", "hlink": "0D5C78", "folHlink": "56636F"}
HEAD, BODY = "Georgia", "Arial"
W, H = 13.333, 7.5
L = 0.6
CW = W - 2 * L

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def I(v):
    return Emu(int(round(v * 914400)))


# ---------------------------------------------------------------- theme and layouts
def set_theme():
    part = prs.slide_master.part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    cs = root.find(".//a:clrScheme", NS)
    cs.set("name", "seqeron")
    for tag, val in COLORS.items():
        el = cs.find("a:" + tag, NS)
        for c in list(el):
            el.remove(c)
        etree.SubElement(el, qn("a:srgbClr")).set("val", val)
    fs = root.find(".//a:fontScheme", NS)
    fs.set("name", "seqeron")
    fs.find("a:majorFont/a:latin", NS).set("typeface", HEAD)
    fs.find("a:minorFont/a:latin", NS).set("typeface", BODY)
    root.find(".//a:themeElements", NS).getparent().set("name", "seqeron")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


if "theme" not in OFF:
    set_theme()
L_TITLE, L_SECTION, L_CONTENT = prs.slide_layouts[0], prs.slide_layouts[2], prs.slide_layouts[5]


def place(ph, x, y, w, h):
    ph.left, ph.top, ph.width, ph.height = I(x), I(y), I(w), I(h)


for lay in (L_TITLE, L_SECTION):
    lay.background.fill.solid()
    lay.background.fill.fore_color.theme_color = DK2
for ph in L_TITLE.placeholders:
    if ph.placeholder_format.idx == 0:
        place(ph, 0.8, 2.0, 11.7, 1.4)
    elif ph.placeholder_format.idx == 1:
        place(ph, 0.8, 3.45, 11.7, 1.0)
for ph in L_SECTION.placeholders:
    if ph.placeholder_format.idx == 0:
        place(ph, 0.8, 3.0, 11.7, 1.1)
    elif ph.placeholder_format.idx == 1:
        place(ph, 0.8, 4.15, 11.7, 0.8)
for ph in L_CONTENT.placeholders:
    if ph.placeholder_format.idx == 0:
        place(ph, L, 0.35, CW, 0.75)

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def footer_sp(sid, name, x, w, para):
    return etree.fromstring(
        f'<p:sp xmlns:p="{P}" xmlns:a="{A}"><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/>'
        f'<p:cNvSpPr txBox="1"/><p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr><a:xfrm>'
        f'<a:off x="{I(x)}" y="{I(7.0)}"/><a:ext cx="{I(w)}" cy="{I(0.3)}"/></a:xfrm>'
        f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr><p:txBody>'
        f'<a:bodyPr wrap="none" lIns="0" tIns="0" rIns="0" bIns="0" anchor="ctr"/><a:lstStyle/>'
        f"{para}</p:txBody></p:sp>")


RPR = '<a:rPr lang="en-US" sz="1000" dirty="0"><a:solidFill><a:schemeClr val="accent5"/></a:solidFill></a:rPr>'
tree = L_CONTENT.shapes._spTree
if "footer" not in OFF: tree.append(footer_sp(901, "Footer", L, 6.0, f"<a:p xmlns:a=\"{A}\"><a:r>{RPR}<a:t>seqeron</a:t></a:r></a:p>"))


# ---------------------------------------------------------------- primitives
_n = [0]


def nm(kind):
    _n[0] += 1
    return f"{kind} {_n[0]}"


def runfmt(r, size=None, bold=None, color=None, italic=None, head=False):
    f = r.font
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.theme_color = color
    if head:
        f.name = "+mj-lt"
        r._r.get_or_add_rPr().set("cap", "none")


def fill_tf(tf, paras, size=14, color=TX, align=PP_ALIGN.LEFT, bold=False, space=0, head=False):
    """paras: str | list of str | list of dict(text|runs, size, bold, color, bullet, align, space, head)."""
    if isinstance(paras, str):
        paras = [paras]
    first = True
    for p in paras:
        if isinstance(p, str):
            p = {"text": p}
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        para.alignment = p.get("align", align)
        sp = p.get("space", space)
        if sp:
            para.space_after = Pt(sp)
        runs = p.get("runs") or [(p.get("text", ""), {})]
        for txt, o in runs:
            r = para.add_run()
            r.text = txt
            runfmt(r, o.get("size", p.get("size", size)), o.get("bold", p.get("bold", bold)),
                   o.get("color", p.get("color", color)), o.get("italic", p.get("italic")),
                   o.get("head", p.get("head", head)))
        if p.get("bullet") and "bullets" not in OFF:
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(I(0.25)))
            pPr.set("indent", str(-I(0.25)))
            bc = etree.SubElement(pPr, qn("a:buClr"))
            etree.SubElement(bc, qn("a:schemeClr")).set("val", "accent1")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")


def text(s, x, y, w, h, paras, size=14, color=TX, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         bold=False, space=0, head=False, margin=0.0):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tb.name = nm("Text")
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, I(margin))
    fill_tf(tf, paras, size, color, align, bold, space, head)
    return tb


STYLES = {
    # fill, line color, line width pt, dash
    "box": (BG, TX, 1.0, None),
    "host": (LT2, A6, 1.25, None),
    "mem": (LT2, A5, 1.0, MSO_LINE.DASH),
    "queue": (A3, A1, 1.25, None),
    "syn": (A4, A2, 1.25, None),
    "held": (A3, A1, 1.25, MSO_LINE.DASH),
    "miss": (None, A2, 1.25, MSO_LINE.DASH),
    "lead": (BG, A1, 2.5, None),
    "bar": (A1, None, 0, None),
    "dim": (A6, None, 0, None),
    "card": (LT2, None, 0, None),
    "dark": (DK2, None, 0, None),
    "accent": (A1, None, 0, None),
    "warm": (A2, None, 0, None),
    "plain": (BG, A6, 1.0, None),
    "frame": (None, A5, 1.0, MSO_LINE.DASH),
    "label": (BG, None, 0, None),
}


def shape(s, x, y, w, h, style="box", kind=MSO_SHAPE.RECTANGLE, name="Shape"):
    fill, line, width, dash = STYLES[style]
    sh = s.shapes.add_shape(kind, I(x), I(y), I(w), I(h))
    sh.name = nm(name)
    sh.shadow.inherit = False
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.theme_color = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.theme_color = line
        sh.line.width = Pt(width)
        if dash:
            sh.line.dash_style = dash
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = 0.08
    return sh


def box(s, x, y, w, h, title=None, sub=None, style="box", tsize=12, ssize=10, kind=MSO_SHAPE.RECTANGLE,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, tcolor=TX, scolor=A5, margin=0.05, name="Box"):
    sh = shape(s, x, y, w, h, style, kind, name)
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, I(margin))
    paras = []
    if title:
        paras.append({"text": title, "size": tsize, "bold": True, "color": tcolor})
    if sub:
        paras.append({"text": sub, "size": ssize, "color": scolor})
    if paras:
        fill_tf(tf, paras, align=align)
    return sh


LINES = {"ln": (TX, 1.25, None), "a": (A1, 1.75, None), "w": (A2, 1.75, None), "d": (A5, 1.25, MSO_LINE.DASH),
         "lt": (A3, 1.25, None)}


def _style_line(ln_shape, style, head, tail_start=False):
    color, width, dash = LINES[style]
    ln_shape.line.color.theme_color = color
    ln_shape.line.width = Pt(width)
    if dash:
        ln_shape.line.dash_style = dash
    ln = ln_shape.line._get_or_add_ln()
    if "heads" in OFF: return
    if tail_start:
        etree.SubElement(ln, qn("a:headEnd")).attrib.update({"type": "triangle", "w": "med", "len": "med"})
    if head:
        etree.SubElement(ln, qn("a:tailEnd")).attrib.update({"type": "triangle", "w": "med", "len": "med"})


def arrow(s, pts, style="ln", head=True, both=False):
    if len(pts) == 2:
        (x1, y1), (x2, y2) = pts
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
        c.name = nm("Arrow")
        _style_line(c, style, head, both)
        return c
    if "freeform" in OFF: return None
    fb = s.shapes.build_freeform(I(pts[0][0]), I(pts[0][1]), scale=1.0)
    fb.add_line_segments([(I(x), I(y)) for x, y in pts[1:]], close=False)
    sh = fb.convert_to_shape()
    sh.name = nm("Path")
    sh.fill.background()
    sh.shadow.inherit = False
    _style_line(sh, style, head, both)
    return sh


def table(s, x, y, colw, rows, size=14, rowh=0.45, head_fill=A1):
    n, m = len(rows), len(colw)
    gt = s.shapes.add_table(n, m, I(x), I(y), I(sum(colw)), I(rowh * n))
    gt.name = nm("Table")
    tbl = gt.table
    tbl.horz_banding = False
    tbl.first_row = True
    for j, w in enumerate(colw):
        tbl.columns[j].width = I(w)
    for i, row in enumerate(rows):
        tbl.rows[i].height = I(rowh)
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = I(0.1)
            c.margin_top = c.margin_bottom = I(0.05)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.theme_color = head_fill
                fill_tf(c.text_frame, [{"text": val, "bold": True, "color": BG, "size": size - 1}])
            else:
                c.fill.fore_color.theme_color = LT2 if i % 2 == 0 else BG
                fill_tf(c.text_frame, [{"text": val, "size": size, "bold": j == 0}])
    return gt


class Fig:
    """Maps a diagram drawn in abstract units onto the slide."""

    def __init__(self, s, ox, oy, sc):
        self.s, self.ox, self.oy, self.sc = s, ox, oy, sc

    def X(self, u):
        return self.ox + u * self.sc

    def Y(self, u):
        return self.oy + u * self.sc

    def box(self, x, y, w, h, title=None, sub=None, style="box", **kw):
        return box(self.s, self.X(x), self.Y(y), w * self.sc, h * self.sc, title, sub, style, **kw)

    def host(self, x, y, w, h, label, align=PP_ALIGN.CENTER, size=11):
        sh = shape(self.s, self.X(x), self.Y(y), w * self.sc, h * self.sc, "host", name="Host")
        tf = sh.text_frame
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_top, tf.margin_left, tf.margin_right = I(0.08), I(0.12), I(0.12)
        fill_tf(tf, [{"text": label.upper(), "size": size, "bold": True, "color": A5, "align": align}])
        return sh

    def label(self, x, y, w, h, t, size=10, color=A5, align=PP_ALIGN.CENTER, bold=False):
        return text(self.s, self.X(x), self.Y(y), w * self.sc, h * self.sc, t, size=size, color=color,
                    align=align, anchor=MSO_ANCHOR.MIDDLE, bold=bold)

    def arrow(self, pts, style="ln", head=True, both=False):
        return arrow(self.s, [(self.X(x), self.Y(y)) for x, y in pts], style, head, both)


def fits(label, size, w):
    """Warns of a one-line label likely to wrap: Arial averages about 0.5 em a character."""
    if len(label) * size * 0.5 / 72 > w:
        print(f"wraps at {w:.2f} in: {label}")


class Seq:
    """A sequence diagram: lanes left to right, rows top to bottom from a cursor."""

    def __init__(self, s, lanes, xs, y=1.58, bottom=6.85, row=0.32,
                 legend=(("live queue", "a"), ("replay 201", "w"), ("request", "ln"), ("reply", "d"))):
        self.s, self.xs, self.bottom, self.row = s, xs, bottom, row
        self.frames = []
        for (name, sub, style), x in zip(lanes, xs):
            box(s, x - 1.05, y, 2.1, 0.46, name, sub, style=style, tsize=12, ssize=9, name="Lane")
        for x in xs:
            arrow(s, [(x, y + 0.46), (x, bottom)], "lt", head=False)
        self.y = y + 0.46
        self.legend(legend)

    def legend(self, items):
        x = W - L - 1.1 * len(items)
        for lab, st in items:
            arrow(self.s, [(x, 1.31), (x + 0.35, 1.31)], st)
            text(self.s, x + 0.42, 1.19, 0.68, 0.24, lab, size=10, color=A5, anchor=MSO_ANCHOR.MIDDLE)
            x += 1.1

    def msg(self, a, b, label, style="ln", lost=False):
        """A message from lane a to lane b, labelled above its line."""
        self.y += self.row
        x1, x2 = self.xs[a], self.xs[b]
        lo, hi = min(x1, x2), max(x1, x2)
        fits(label, 11, hi - lo - 0.1)
        text(self.s, lo + 0.05, self.y - 0.27, hi - lo - 0.1, 0.25, label, size=11,
             color=A2 if lost else TX, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)
        if lost:
            xe = x1 + (x2 - x1) * 0.62
            arrow(self.s, [(x1, self.y), (xe, self.y)], "w", head=False)
            text(self.s, xe - 0.1, self.y - 0.14, 0.2, 0.28, "×", size=16, bold=True, color=A2,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        else:
            arrow(self.s, [(x1, self.y), (x2, self.y)], style)

    def note(self, lane, label, w=4.1, style="card"):
        """A step a lane takes on its own, to the right of its lifeline."""
        self.y += self.row
        fits(label, 11, w - 0.16)
        box(self.s, self.xs[lane] + 0.08, self.y - 0.27, w, 0.3, None, None, style=style, name="Note")
        text(self.s, self.xs[lane] + 0.16, self.y - 0.27, w - 0.16, 0.3, label, size=11, anchor=MSO_ANCHOR.MIDDLE)

    def _tag(self, tag, cond, x, y):
        box(self.s, x, y, 0.62, 0.24, tag, None, style="dim", tsize=9, margin=0, name="Frame tag")
        if cond:
            shape(self.s, x + 0.66, y, 0.12 + 0.072 * len(cond), 0.24, "label", name="Frame guard")
            text(self.s, x + 0.72, y, 0.072 * len(cond), 0.24, [{"text": cond, "italic": True}], size=10,
                 color=A5, anchor=MSO_ANCHOR.MIDDLE).text_frame.word_wrap = False

    def begin(self, tag, cond, x=L, w=CW):
        """Opens an alt or loop frame; the guard sits in its top row."""
        self.y += 0.12
        self.frames.append((x, w, self.y))
        self._tag(tag, cond, x, self.y)
        self.y += 0.15

    def orelse(self, cond):
        x, w, _ = self.frames[-1]
        self.y += 0.14
        arrow(self.s, [(x, self.y), (x + w, self.y)], "d", head=False)
        self._tag("ELSE", cond, x, self.y)
        self.y += 0.15

    def end(self):
        x, w, top = self.frames.pop()
        self.y += 0.12
        shape(self.s, x, top, w, self.y - top, "frame", name="Frame")


# ---------------------------------------------------------------- slide kinds
SECTIONS = []  # (name, [slide ids])


def section(name):
    SECTIONS.append((name, []))


def _track(s):
    SECTIONS[-1][1].append(s.slide_id)


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


def content(title, kicker=None, note=""):
    s = prs.slides.add_slide(L_CONTENT)
    _track(s)
    t = s.shapes.title
    tf = t.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    fill_tf(tf, [{"text": title, "size": 34, "bold": True, "color": TX, "head": True}])
    if kicker:
        text(s, L, 1.08, CW, 0.45, kicker, size=16, color=A5)
    s.shapes._spTree.append(footer_sp(9000 + len(prs.slides), "Slide number", W - L - 1.0, 1.0,
                            f'<a:p xmlns:a="{A}"><a:pPr algn="r"/><a:fld id="{{{uuid.uuid4()}}}" type="slidenum">'
                            f'{RPR}<a:t>{len(prs.slides)}</a:t></a:fld></a:p>'))
    notes(s, note)
    return s


def tape(s, x, y, n, current, cw=0.42, gap=0.08, on_dark=True):
    """The deck's motif: a run of globalSeqNo cells, the current one filled."""
    for i in range(n):
        st = "accent" if i + 1 == current else ("dark" if on_dark else "card")
        sh = shape(s, x + i * (cw + gap), y, cw, cw, st, name="Tape cell")
        if on_dark and i + 1 != current:
            sh.line.color.theme_color = A1
            sh.line.width = Pt(1.25)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        fill_tf(tf, [{"text": str(i + 1), "size": 12, "bold": i + 1 == current,
                      "color": BG if i + 1 == current else A3, "align": PP_ALIGN.CENTER}])


def log_cells(s, x, y, w=0.31, h=0.36, step=0.346):
    """A queue as a log: frames 1-4 written, 5 being written, two free cells."""
    for i in range(7):
        st = "accent" if i == 4 else ("queue" if i < 4 else "plain")
        sh = shape(s, x + i * step, y, w, h, st, name="Queue cell")
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        if i < 5:
            fill_tf(tf, [{"text": str(i + 1), "size": 9, "bold": i == 4, "color": BG if i == 4 else A1,
                          "align": PP_ALIGN.CENTER}])


def num_circle(s, x, y, n, d=0.5, style="accent", size=16):
    sh = shape(s, x, y, d, d, style, MSO_SHAPE.OVAL, "Number")
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    fill_tf(tf, [{"text": str(n), "size": size, "bold": True, "color": BG, "align": PP_ALIGN.CENTER}])
    return sh


def card(s, x, y, w, h, paras, style="card", margin=0.2, anchor=MSO_ANCHOR.TOP):
    sh = shape(s, x, y, w, h, style, name="Card")
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, I(margin))
    fill_tf(tf, paras)
    return sh


C, M, R = PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, PP_ALIGN.RIGHT

# ================================================================= OPENING
section("Opening")
s = prs.slides.add_slide(L_TITLE)
_track(s)
for ph in s.placeholders:
    tf = ph.text_frame
    tf.margin_left = tf.margin_right = 0
    if ph.placeholder_format.idx == 0:
        fill_tf(tf, [{"text": "seqeron", "size": 66, "bold": True, "color": BG, "head": True}])
    else:
        fill_tf(tf, [{"text": "A replicated sequencer on Aeron Cluster", "size": 26, "color": A3}])
tape(s, 0.8, 5.2, 12, 7, cw=0.52, gap=0.08)
text(s, 0.8, 5.85, 11.7, 0.4, "Architecture briefing  ·  seqeron 0.13.0  ·  Aeron 1.53.2", size=14, color=A3)
url = text(s, 0.8, 6.25, 11.7, 0.4, "github.com/FredrikJDahlberg/seqeron", size=14, color=A3)
url.click_action.hyperlink.address ="https://github.com/FredrikJDahlberg/seqeron"
notes(s, "seqeron is the sequencing tier: an Aeron Cluster replicated state machine that gives messages from any "
         "number of producers one global, gap-free total order, plus the replay service and client library that "
         "deliver that order to every consumer. The row of cells is the deck's motif: the log, numbered by "
         "globalSeqNo.")

s = content("Summary", "What seqeron is, what is proven, what a client gets, and how to try it",
            "The one-minute version. seqeron gives every producer's messages one global, gap-free order and every "
            "consumer the same complete history on its own host. Its correctness properties are built in and tested "
            "against a live cluster. Confirmed ingress is exactly once while the producer runs; a restarted producer "
            "starts with nothing pending, and leader-only replies are at least once across a failover, deduplicated "
            "by key. A client programs against one façade per kind and writes only its own logic: the "
            "session, recovery, confirmation and failure handling are behind it, the same in Java, C++ and C#. It can be "
            "evaluated hands-on: a release runs a one-member cluster, the examples cover each kind of client and "
            "measure a round trip on your own hardware, and the failure harnesses run against a live cluster. What a "
            "production deployment needs beyond that, performance figures, a bounded archive, wire security and "
            "disaster recovery, is in the appendix.")
text(s, L, 1.75, CW, 0.8, "A Raft-replicated sequencer on Aeron Cluster: one global, gap-free order for every "
     "producer's messages, and the same complete history on every consumer's host. The cluster is Java; clients are "
     "Java, C++ and C#. Open source, Apache 2.0.", size=17)
summary = [("Proven", A1, [[("Determinism: every decision is a function of the log (S-3)", {})],
                           [("A minority of members can crash with no gap and no renumbering", {})],
                           [("Confirmed ingress: exactly once, in order, while its producer runs", {})],
                           [("A live-cluster harness for every failure in the model, run in CI", {})]]),
           ("What a client gets", A2, [[("One façade, one call per duty cycle, on the caller's thread", {})],
                                       [("History then live, gap-free, from its own host", {})],
                                       [("Its own messages confirmed, and resent across a failover", {})],
                                       [("An elected standby or a leader gate, decided from the log", {})]]),
           ("How to evaluate it", TX, [[("A release runs a one-member cluster; getting-started.md adds a consumer "
                                         "and a producer", {})],
                                       [("Examples of each kind in Java, C++ and C#; a snapshot restore in Java "
                                         "and C++", {})],
                                       [("The examples ping through the cluster: a round trip measured on your "
                                         "own hardware", {})],
                                       [("The failure harnesses: leader kills, chaos, a containerized soak", {})]])]
for i, (t, c, pts) in enumerate(summary):
    x = L + i * (3.84 + 0.305)
    card(s, x, 2.9, 3.84, 3.5, [{"text": t, "size": 18, "bold": True, "head": True, "color": c, "space": 8}]
         + [{"runs": p, "size": 14, "bullet": True, "space": 6} for p in pts], margin=0.22)
text(s, L, 6.55, CW, 0.4, [{"runs": [("Not yet established: ", {"bold": True, "color": A2}),
                                    ("performance figures, a bounded archive, wire security, disaster recovery. "
                                     "The appendix covers each.", {})]}], size=14)

s = content("One replicated order",
            "Every input is stamped once; every service is a deterministic state machine over the result",
            "The sequencer architecture: services never coordinate with each other. Two replicas that apply the same "
            "stream hold the same state; a hot standby is a replica that does not publish; an audit is a replay; a "
            "restart is a replay from a known point. The sequencer then becomes the component everything depends on, "
            "so seqeron makes it a Raft-replicated state machine that survives the loss of a minority of nodes and "
            "continues under a new leader with no gap and no renumbering. Payloads are opaque: the cluster copies "
            "them through unopened.")
for i, t in enumerate(["Gateway pairs", "Applications", "Operator (clusterctl)"]):
    box(s, L, 1.95 + i * 0.8, 2.5, 0.6, t, style="box", tsize=13)
    arrow(s, [(L + 2.5, 2.25 + i * 0.8), (3.65, 3.05)], "ln")
box(s, 3.7, 2.15, 2.6, 1.8, "Sequencer cluster", "Raft · 3 to 7 cluster members", style="lead", tsize=15, ssize=12)
arrow(s, [(6.3, 3.05), (6.75, 3.05)], "a")
for i in range(6):
    box(s, 6.8 + i * 0.6, 2.78, 0.54, 0.54, str(101 + i), style="queue", tsize=12)
text(s, 6.8, 2.25, 3.6, 0.4, "globalSeqNo: no gaps, no reuse", size=12, color=A5)
arrow(s, [(10.36, 3.05), (10.85, 3.05)], "a")
card(s, 10.9, 1.95, 1.83, 2.2, [{"text": "Every consumer", "bold": True, "size": 13},
                                 {"text": "applications", "size": 13, "bullet": True},
                                 {"text": "gateways", "size": 13, "bullet": True},
                                 {"text": "audit, risk", "size": 13, "bullet": True}], style="plain", margin=0.12)
stats = [("1", "stream", "one total order across every producer"),
         ("3–7", "members", "a minority can fail with no gap and no renumbering"),
         ("44 B", "envelope", "on a sequenced message; payload never decoded"),
         ("1 Hz", "cluster clock", "a heartbeat on consensus time every client agrees on")]
for i, (big, lab, d) in enumerate(stats):
    x = L + i * (2.85 + 0.243)
    card(s, x, 4.75, 2.85, 1.55, [{"runs": [(big, {"size": 36, "bold": True, "color": A1, "head": True}),
                                            ("  " + lab, {"size": 14, "bold": True, "color": TX})]},
                                  {"text": d, "size": 13, "color": A5}], margin=0.18)

s = content("Context and alternatives", "The problem it solves, and the options it is weighed against",
            "The alternatives are compared on properties. Kafka on one partition gives a total order and mature "
            "tooling, but no clock that advances while producers are silent and no active/standby producer election. "
            "Aeron without clustering gives reliable UDP and IPC transport with recording and replay, but no "
            "replication: the sequencer is one process and a single point of failure. Aeron Cluster used directly "
            "gives the same consensus and log, but each cluster then rebuilds the queue, replay, the election, confirmed "
            "ingress and the fences; the slide Aeron Cluster used directly is that list. A database sequence is familiar and "
            "transactional, but sequences skip values on rollback, so the order is not gap-free, and every message is "
            "a database write.")
card(s, L, 1.75, CW, 1.1, [
    {"text": "Services that must agree on what happened, and in what order: a trading venue, an order management "
             "system, a risk engine. The sequencer architecture settles the order once, in one place, and every "
             "service becomes a deterministic replica of one stream.", "size": 15}], margin=0.2)
table(s, L, 3.1, [2.6, 4.4, 5.13], [
    ["Option", "Gives", "Drawback"],
    ["Kafka, one partition", "Total order, retention, mature tooling", "No cluster clock or producer election"],
    ["Aeron", "Reliable UDP and IPC; recording, replay", "No replication: one sequencer, one point of failure"],
    ["Aeron Cluster alone", "The same consensus and log", "Replay, election and confirmed ingress rebuilt per cluster"],
    ["Database sequence", "Familiar and transactional", "Skips values on rollback; a write per message"],
], size=14, rowh=0.55)
text(s, L, 6.1, CW, 0.4, [{"runs": [("A total order is one serialization point ", {"bold": True}),
                                     ("in every option, seqeron's included: its throughput ceiling is the same in "
                                      "kind.", {})]}], size=14, color=A5)

# ================================================================= AERON
section("Aeron and Aeron Cluster")
s = content("Aeron Cluster", "Raft over Aeron's transport and archive: replicated log, deterministic services",
            "Each member runs a consensus module, an archive and a clustered service container. A client opens a "
            "session and sends on ingress to the leader; responses come back on egress. The leader appends each message "
            "to the Raft log, a stream every member's archive records; followers report append positions and an entry "
            "commits at a majority. Every cluster member's service consumes committed entries in the same order. Session "
            "events, timers and the leader's timestamp travel in the log, so time is deterministic too. A restarting "
            "member normally loads a snapshot and replays the log after it; seqeron disables snapshots, as the Decisions slide in the appendix covers. Aeron's transport model, media driver and archive are in the appendix.")
f = Fig(s, L, 1.75, CW / 960)
f.host(10, 120, 170, 110, "Cluster client")
f.label(10, 160, 170, 24, "AeronCluster", size=12, color=TX, bold=True)
f.label(10, 186, 170, 24, "one session")
f.host(240, 20, 330, 340, "Leader", align=PP_ALIGN.LEFT)
f.box(260, 60, 290, 60, "Consensus module", "election · append · commit · cluster time")
f.box(260, 150, 130, 60, "Archive", "records the log")
f.box(410, 150, 140, 60, "Clustered service", "applies committed log")
f.arrow([(325, 120), (325, 150)])
f.arrow([(480, 120), (480, 150)])
f.label(258, 222, 250, 22, "commit = a majority has appended", align=PP_ALIGN.LEFT)
f.label(258, 244, 250, 34, "timers, session events and timestamps ride the log", align=PP_ALIGN.LEFT)
f.arrow([(180, 160), (260, 95)], "a")
f.label(150, 112, 80, 20, "ingress", color=A1)
f.arrow([(530, 210), (530, 300), (95, 300), (95, 230)], "d")
f.label(150, 304, 300, 22, "egress: responses and session events")
for oy in (20, 200):
    f.host(630, oy, 320, 160, "Follower", align=PP_ALIGN.LEFT)
    f.box(650, oy + 35, 130, 46, "Consensus module")
    f.box(800, oy + 35, 130, 46, "Archive")
    f.box(650, oy + 95, 280, 46, "Clustered service", "same input, same state")
f.arrow([(550, 82), (650, 78)], "a")
f.arrow([(550, 96), (650, 258)], "a")
f.label(560, 52, 80, 20, "log", color=A1)
text(s, L, 6.4, CW, 0.4, [{"runs": [("Member: ", {"bold": True}),
                                    ("one node of the cluster: consensus module, archive, clustered service. In this deck, "
                                     "always a cluster member.", {})]}], size=14)

s = content("What seqeron takes from Aeron", "Three properties it builds on rather than reimplements",
            "Point to land: everything in the rest of the talk is built from these three. Explicit back-pressure lets "
            "the queue be reliable but bounded. Local recording with replay by position makes every member a complete "
            "history server. A replicated log with deterministic time makes the order, the clock and the gateway "
            "election the same on every node. Aeron, Agrona and SBE are open source under Apache 2.0.")
cols = [("Explicit back-pressure", "Offers fail visibly instead of queuing.",
         "A reliable queue that is still bounded; untethered consumers that cannot slow the sequencer."),
        ("Local recording, replay by position", "Each member's archive records any stream to disk, read back from any byte.",
         "Every cluster member keeps full history; every client recovers from its own host."),
        ("Replicated log, deterministic time", "Raft order plus the leader's timestamp in the log.",
         "One total order, a cluster clock, and an election every node decides identically.")]
for i, (t, a, b) in enumerate(cols):
    x = L + i * (3.84 + 0.305)
    shape(s, x, 1.85, 3.84, 4.35, "card", name="Column")
    num_circle(s, x + 0.25, 2.1, i + 1)
    text(s, x + 0.25, 2.8, 3.34, 0.9, t, size=18, bold=True, head=True)
    text(s, x + 0.25, 3.75, 3.34, 0.3, "AERON GIVES", size=11, bold=True, color=A5)
    text(s, x + 0.25, 4.05, 3.34, 0.8, a, size=14)
    text(s, x + 0.25, 4.95, 3.34, 0.3, "SEQERON BUILDS", size=11, bold=True, color=A1)
    text(s, x + 0.25, 5.25, 3.34, 0.9, b, size=14)
text(s, L, 6.35, CW, 0.4, [{"runs": [("Open source: ", {"bold": True, "color": A1}),
                                    ("Aeron, Agrona and SBE, all Apache 2.0.", {})]}], size=14)

s = content("Aeron Cluster used directly", "What every cluster would otherwise build for itself",
            "This is the case for seqeron to an Aeron user: what a team building directly on Aeron Cluster writes "
            "itself, row by row. An offer result says the frame reached the leader's ingress, nothing more, and a "
            "failover loses whatever the old leader had not committed while the session survives. The archive replays "
            "bytes from a position and knows nothing of gaps or of catching up, so joining the live stream from history "
            "without a gap is a state machine to design and prove: ReplayerRecovery, with seeded property tests, in each "
            "language. Cluster timers exist only inside the clustered service, so a client has no clock while producers "
            "are silent. Aeron has no notion of an active/standby producer pair, and no way for a client to know its "
            "view of the log is stale. Its cluster clients are Java and, trailing it, Aeron.NET; the C++ one here is "
            "written against the wire protocol, and C# runs on Aeron.NET's with every timeout set explicitly.")
table(s, L, 1.8, [2.85, 4.4, 4.88], [
    ["Need", "Aeron Cluster gives", "seqeron adds"],
    ["Know a message is sequenced", "An offer result; failover loses uncommitted ingress silently",
     "Confirmed ingress: held until its own queue shows it, resent after a term change"],
    ["Join from history", "Archive replay by byte position; no gaps, no catching up",
     "Replay-to-live with gap repair and an anchor check, property-tested"],
    ["A clock while producers are silent", "Timers inside the clustered service only",
     "ClusterHeartbeat: 1 Hz on consensus time, on every queue"],
    ["Active/standby producers", "Nothing", "The gateway election, decided from the log"],
    ["Services outside the member", "Clustered services, in the member's JVM",
     "Applications on a member or remote, with a leader gate"],
    ["Clients beyond Java", "A Java cluster client; Aeron.NET trails Aeron",
     "C++ cluster client (wire protocol) and C# façade"],
    ["Stop on a stale view", "Nothing", "Fences: close the edge, release the session, report once"],
], size=13, rowh=0.56)
card(s, L, 6.45, CW, 0.5, [{"runs": [("Aeron provides the mechanisms ", {"bold": True}),
                                     ("and almost none of the policy: ", {}),
                                     ("confirmation, history, recovery, failure handling.", {"bold": True})]}],
     style="plain", margin=0.12, anchor=MSO_ANCHOR.MIDDLE)


# ================================================================= GLANCE
section("seqeron overview")
s = content("A member and a client", None,
            "Each member is one JVM: consensus module, sequencer, embedded media driver, archive and Replayer. Live data "
            "reaches a client straight off the queue, untethered. History comes from the member's own recording through "
            "the Replayer on stream 201, with requests on 202 and control replies on 203. Ingress is IPC when this "
            "member leads and UDP otherwise. Other members run the same pipeline and produce a byte-identical queue. The client tier is Java, C++ and C#, "
            "with C# deployed on Windows too; the member is Java only.")
sc = 5.25 / 440
f = Fig(s, (W - 960 * sc) / 2, 1.5, sc)
f.host(10, 60, 170, 90, "Other members")
f.label(10, 95, 170, 40, "same pipeline, byte-identical queue", size=11)
f.arrow([(95, 60), (95, 22), (590, 22), (590, 80)], "d", both=True)
f.label(240, 0, 200, 20, "Raft log replication")
f.host(210, 40, 510, 380, "Cluster member · one JVM", align=PP_ALIGN.LEFT)
f.box(230, 80, 200, 56, "Sequencer", "number · validate · synthesize")
f.box(480, 80, 220, 56, "Consensus module", "Raft log · cluster time")
f.arrow([(480, 108), (430, 108)])
f.label(230, 138, 170, 22, "Queue · aeron:ipc 205", size=10, color=A1, align=PP_ALIGN.LEFT, bold=True)
log_cells(s, f.X(230), f.Y(162), 26 * sc, 30 * sc, 29 * sc)
f.arrow([(359, 136), (359, 162)])
f.box(230, 270, 200, 56, "Archive", "records the queue · tethered")
f.arrow([(301, 192), (301, 270)])
f.box(480, 270, 220, 56, "ReplayerService", "history from the recording", "syn")
f.arrow([(430, 298), (480, 298)])
f.box(230, 350, 470, 50, "Embedded media driver", "queue 205 · replay 201 · request 202 · control 203", "mem")
f.host(760, 60, 190, 290, "Client process")
f.box(772, 100, 166, 56, "ClusterStreamSender", "+ PendingSends", tsize=11)
f.box(772, 190, 166, 56, "ReplayerStreamReceiver", "+ ReplayerRecovery", tsize=11)
f.label(760, 280, 190, 46, "inside the Gateway or Application façade", size=10)
f.label(760, 318, 190, 24, "Java · C++ · C#", size=10, bold=True)
f.arrow([(772, 120), (700, 120)])
f.label(700, 96, 72, 22, "ingress")
f.arrow([(330, 192), (330, 212), (772, 212)], "a")
f.label(480, 188, 220, 22, "live · untethered", color=A1)
f.arrow([(700, 292), (772, 240)], "w")
f.label(712, 296, 60, 22, "replay", color=A2)

s = content("The life of a message", "Six steps from a producer's offer to its own confirmation",
            "1. IngressPublisher wraps the producer's own encoding in an Unsequenced frame and offers it. 2. The leader "
            "appends it to the Raft log; it commits at a majority. 3. Every cluster member's sequencer validates it and assigns "
            "the next globalSeqNo and consensus timestamp. 4. Every archive records it. 5. Each client reads it off its "
            "local queue, checks continuity, dispatches. 6. The producer sees its own frame on its queue; only then is the "
            "message known to be sequenced. A cold start runs steps 5 and 6 against a replay first.")
steps = [("Offer", "Producer wraps its payload in a 28-byte-overhead frame."),
         ("Commit", "Leader appends to the Raft log; a majority holds it."),
         ("Sequence", "Every cluster member numbers, stamps, publishes on its queue."),
         ("Record", "Every cluster member's archive records the queue."),
         ("Deliver", "Each client reads its local queue, checks continuity."),
         ("Confirm", "The producer sees its own frame. Now it is sequenced.")]
sw, sg = 1.86, 0.194
arrow(s, [(L + 0.5, 2.25), (L + 5 * (sw + sg) + 0.5, 2.25)], "lt", head=False)
for i, (t, d) in enumerate(steps):
    x = L + i * (sw + sg)
    num_circle(s, x + 0.25, 2.0, i + 1, style="warm" if i == 5 else "accent")
    text(s, x, 2.75, sw, 0.45, t, size=18, bold=True, head=True)
    text(s, x, 3.25, sw - 0.1, 1.9, d, size=14)
card(s, L, 5.35, CW, 1.25, [{"runs": [("A send that succeeds is not a message sequenced. ", {"bold": True}),
                                      ("The cluster confirms nothing on egress; step 6 is the producer's own confirmation, "
                                       "on its own queue.", {})], "size": 16}], margin=0.25, anchor=M)

s = content("The queue", "Every cluster member publishes and records its own copy; none is replicated",
            "All members apply the same committed log in the same order with the same state, so their queues are "
            "byte-identical frame for frame. Each archive therefore holds complete history: no cross-node replication of "
            "recordings and no leader-only archive. The queue is created once at start and never re-created on a "
            "leadership change, so a recording is one continuous run across every leader tenure. The offer is reliable "
            "because a dropped frame is a permanent hole, and bounded because only the archive, the one tethered "
            "subscriber, can hold it back. A slow archive that still advances is never timed out. One that makes no "
            "progress for 200 ms raises the stall gauge, and after 1 s the member exits with code 70 rather than "
            "throw: a throwing cluster callback has already had its log position advanced.")
for i in range(3):
    x = L + i * 1.95
    text(s, x, 1.85, 1.75, 0.35, f"member {i}", size=12, bold=True, color=A5, align=C)
    box(s, x, 2.25, 1.75, 0.6, "Raft log", style="box", tsize=13)
    arrow(s, [(x + 1.11, 2.85), (x + 1.11, 3.46)])
    text(s, x, 3.22, 1.0, 0.24, "Queue 205", size=10, bold=True, color=A1, anchor=MSO_ANCHOR.MIDDLE)
    log_cells(s, x, 3.46, 0.22, 0.3, 0.25)
    arrow(s, [(x + 0.61, 3.76), (x + 0.61, 4.15)])
    box(s, x, 4.15, 1.75, 0.6, "Archive", style="box", tsize=13)
text(s, L, 4.95, 5.65, 1.6, [{"text": "Same committed log → byte-identical queues → every archive holds full history.",
                              "bold": True, "size": 15, "space": 6},
                             {"text": "The queue survives leadership changes: one continuous recording per member.", "size": 14, "color": A5}])
text(s, 6.9, 1.85, 5.8, 0.35, "RELIABLE, AND BOUNDED", size=12, bold=True, color=A1)
table(s, 6.9, 2.25, [2.2, 3.63], [
    ["Archive state", "Response"],
    ["Slow, advancing", "Keep retrying; never timed out"],
    ["No progress 200 ms", "Raise the stall gauge"],
    ["No progress 1 s", "Terminate the member (exit 70)"],
    ["Recording gone", "Terminate at once"],
    ["Stopped silently", "Caught by the 1 Hz liveness check"],
], size=14, rowh=0.5)
text(s, 6.9, 5.4, 5.8, 1.2, "A member that cannot keep history stops instead of sequencing history it cannot serve. "
                            "Peers keep quorum; its restart rebuilds the recording.", size=14, color=A5)

s = content("Live from queue and history from archive", "The Replayer serves one recording to many clients",
            "Consumers read the queue directly and untethered: one that falls a window behind is dropped by its driver "
            "within about 0.2 s and heals by replay. History for a cold start or a gap is served by the co-located "
            "ReplayerService; nothing reads the archive except through it. It serves only the active recording, which "
            "holds the whole log because every member start replays from globalSeqNo 1. The Replayer checks nothing "
            "itself: each client checks that its first frame is globalSeqNo 1 and that a resume opens on the frame it "
            "asked for, and stops otherwise. Four replays at once; a slot frees when its replay's publication leaves the "
            "driver, at once when the client closes its replay subscription, within the driver's 10 s liveness timeout "
            "when the client died, below the client's 60 s recovery-stall fence. While the queue's recording trails the "
            "queue by more than a quarter of the publication "
            "window, a new replay starts only when none is running: replays compete with the recorder for the disk, and "
            "the recorder is what holds the sequencer back. seqeron_replayer_throttled_requests_total counts the requests "
            "held for that reason, and a member warns at start on Linux when its Aeron directory is not on tmpfs. Control "
            "replies are retried briefly and dropped, never blocked on.")
box(s, 1.64, 1.6, 1.0, 0.32, "Sequencer", style="box", tsize=11)
arrow(s, [(2.139, 1.92), (2.139, 2.25)])
text(s, L, 1.98, 1.9, 0.24, "Queue · aeron:ipc 205", size=10, bold=True, color=A1, anchor=MSO_ANCHOR.MIDDLE)
log_cells(s, L, 2.25)
box(s, L, 3.6, 2.3, 0.7, "Archive", style="box", tsize=14)
box(s, L, 5.2, 2.3, 0.7, "ReplayerService", style="syn", tsize=14)
arrow(s, [(1.101, 2.61), (1.101, 3.6)])
arrow(s, [(L + 1.15, 4.3), (L + 1.15, 5.2)])
box(s, 3.85, 3.6, 1.9, 0.7, "Client", style="box", tsize=14)
arrow(s, [(1.793, 2.61), (1.793, 2.9), (4.8, 2.9), (4.8, 3.6)], "a")
text(s, 3.1, 2.52, 2.0, 0.35, "live, untethered", size=12, color=A1)
arrow(s, [(L + 2.3, 5.55), (4.8, 5.55), (4.8, 4.3)], "w")
text(s, 3.0, 5.6, 2.0, 0.35, "history", size=12, color=A2)
rows = [("One recording", "The active one holds the whole log; older ones are prefixes nothing reads."),
        ("Checked by its clients", "First frame must be globalSeqNo 1, a resume must open where asked, or the client stops."),
        ("Four slots, one while the recorder is behind", "A slot frees with its replay's publication: a dead client's within 10 s. Clients fence at 60 s."),
        ("Never held by a client", "Control replies get a bounded retry, then drop; clients resend.")]
for i, (t, d) in enumerate(rows):
    y = 1.9 + i * 1.2
    num_circle(s, 6.4, y + 0.05, i + 1, d=0.45, size=14)
    text(s, 7.05, y, 5.68, 0.4, t, size=16, bold=True)
    text(s, 7.05, y + 0.42, 5.68, 0.7, d, size=14, color=A5)

s = content("Header", "Take 18 bytes from ingress, append 16, never touch the payload",
            "Schema 210, two families. The application family carries an opaque payload named by payloadId: Unsequenced "
            "on ingress, Sequenced on the queue. The system family is seqeron's own vocabulary, named by systemEventType "
            "at the same offset 16. Every field is at the same offset in all eight templates, so a consumer tracks "
            "continuity without branching on message type. sessionId is overwritten by the sequencer with the session "
            "the frame arrived on; the other ingress fields are copied. Consumers split by "
            "family first, then dispatch on (payloadId, templateId), never templateId alone. On the wire the 8-byte SBE "
            "message header comes first and the payload's 2-byte length prefix after the header, so a frame carries "
            "8 + 18 + 2 = 28 bytes besides its payload on ingress, and 8 + 34 + 2 = 44 on the queue.")
sc = CW / 880
f = Fig(s, L - 30 * sc + 0.0, 1.75, sc)
for u, lab in [(30, "0"), (126, "4"), (222, "8"), (414, "16"), (462, "18"), (654, "26"), (846, "34")]:
    f.label(u - 20, 10, 40, 22, lab, size=11, color=A5)
f.box(30, 40, 96, 48, "sourceId", tsize=12)
f.box(126, 40, 96, 48, "connectionId", tsize=12)
f.box(222, 40, 192, 48, "sessionId", "set by the sequencer", "syn")
f.box(414, 40, 48, 48, "payloadId", tsize=8, margin=0)
f.box(462, 40, 192, 48, "globalSeqNo", style="queue", tsize=13)
f.box(654, 40, 192, 48, "timestamp", "Raft consensus time, ns", "queue")
f.arrow([(30, 98), (30, 106), (460, 106), (460, 98)], "ln", head=False)
f.label(30, 110, 430, 26, "from ingress · 18 bytes, sessionId rewritten", size=12, color=TX)
f.arrow([(464, 98), (464, 106), (846, 106), (846, 98)], "a", head=False)
f.label(464, 110, 382, 26, "appended by the sequencer · 16 bytes", size=12, color=A1)
f.label(30, 140, 816, 24, "Offset 16 is payloadId in the application family and systemEventType in the system family.",
        size=12)
card(s, L, 4.15, 5.9, 1.95, [{"text": "Application family", "bold": True, "size": 18, "head": True, "space": 6},
                            {"text": "Unsequenced (100) → Sequenced (101)", "size": 14, "bold": True, "space": 6},
                            {"text": "One opaque payload named by payloadId. The cluster decodes no application payload.", "size": 14}])
card(s, L + 6.23, 4.15, 5.9, 1.95, [{"text": "System family", "bold": True, "size": 18, "head": True, "space": 6},
                                   {"text": "UnsequencedSystem (102) → SequencedSystem (103), plus 104–107", "size": 14, "bold": True, "space": 6},
                                   {"text": "Named by systemEventType. Twelve events submitted, four synthesized.", "size": 14}])

s = content("Repairing a gap", "Replay only what was lost; hold live frames meanwhile",
            "A live frame ahead of the next expected number clears isCaughtUp and requests a resume at the position where "
            "the last delivered frame ended. The resumed replay's first frame must be the expected number; if not, the "
            "recording does not match the positions the client holds, and the client stops. Live frames "
            "past the gap go into a bounded FIFO, 65,536 frames or 16 MiB, and drain when the replay reaches them, so "
            "recovery ends without a second gap; past either bound it drops them and resumes again at its last delivered "
            "frame. A replay with no progress for 5 s is requested again. A cold start is "
            "the same walk from the start. All of it is ReplayerRecovery, a pure state machine with case-for-case tests "
            "and seeded property tests in Java, C++ and C#.")
sc = CW / 760
f = Fig(s, L - 10 * sc, 1.75, sc)
for lab, y in (("Queue 205", 40), ("Replay 201", 110), ("Delivered", 180)):
    f.label(10, y, 140, 34, lab, size=14, color=TX, align=PP_ALIGN.LEFT, bold=True)
queuest = {105: "miss", 106: "held", 107: "held", 108: "held"}
for k in range(10):
    n, x = 101 + k, 160 + k * 58
    st = queuest.get(n, "queue")
    f.box(x, 40, 50, 34, "×" if n == 105 else str(n), style=st, tsize=13, tcolor=A2 if n == 105 else TX)
    f.box(x, 180, 50, 34, str(n), style="syn" if n == 105 else "queue", tsize=13)
f.box(392, 110, 50, 34, "105", style="syn", tsize=13)
f.arrow([(417, 76), (417, 108)], "d")
f.label(450, 6, 300, 30, "106 arrives: gap → resume after 104", size=12, color=A2, align=PP_ALIGN.LEFT)
f.label(450, 80, 300, 26, "106–108 retained in a bounded FIFO", size=12, color=A1, align=PP_ALIGN.LEFT)
f.label(450, 112, 300, 30, "anchor check: first frame must be 105", size=12, align=PP_ALIGN.LEFT)
f.label(160, 222, 590, 28, "105 from the replay, then the FIFO drains, then back on the live queue.", size=12,
        align=PP_ALIGN.LEFT)
text(s, L, 6.0, CW, 0.7, [{"runs": [("Cold start ", {"bold": True}),
                                    ("is the same walk from globalSeqNo 1 to the reported end, then the queue, "
                                     "de-duplicating on globalSeqNo.", {})]}], size=14)

# ================================================================= PRODUCERS
section("Gateways and applications")
s = content("Two kinds of producer", "A façade per kind; the application writes its logic and none of seqeron's",
            "A gateway is an edge producer deployed as an active/hot-standby pair named in the topology; the cluster "
            "designates the active instance. An application runs one instance on every member, and a member can host "
            "several applications; LeadershipChanged already picks one instance of each, so only the one on the leader's "
            "host publishes. Or it runs remote, as the one instance on its own "
            "host: its gate opens once caught up, whoever leads, and it submits over UDP. Nothing elects "
            "between two remote instances, so a remote service that needs a standby is a gateway. Both façades run one single-threaded "
            "duty cycle on the caller's thread: doWork() returns a work count for the caller's idle strategy. The client "
            "tier starts no threads, and dispatch and confirmed publish allocate nothing once warm, which the C# suite "
            "tests; Java and C++ have no such test. The deployment is "
            "declared in one XML topology document that clusterctl load-topology puts on the log.")
for i, (t, sub, pts) in enumerate([
        ("Gateway", "active / hot-standby pair",
         ["Exchange session, client connection, feed handler", "The cluster designates the active instance",
          "Owns its external connections and a connection-id space"]),
        ("Application", "one instance on every member, or remote",
         ["Deterministic service over the stream; several per member", "On a member: publishes only while that member leads",
          "Remote: one instance on its own host, ingress over UDP",
          "Owns leader-only work: replies, outbound effects"])]):
    x = L + i * (5.9 + 0.33)
    shape(s, x, 1.85, 5.9, 3.0, "card", name="Producer")
    text(s, x + 0.25, 2.0, 5.4, 0.5, t, size=22, bold=True, head=True)
    text(s, x + 0.25, 2.5, 5.4, 0.35, sub, size=13, color=A1, bold=True)
    text(s, x + 0.25, 3.0, 5.4, 1.8, [{"text": p, "bullet": True, "space": 6} for p in pts], size=14)
text(s, L, 5.1, CW, 0.35, "EACH FAÇADE'S doWork(), ON THE CALLER'S THREAD", size=12, bold=True, color=A5)
chips = ["Cluster session", "Queue + recovery", "Confirmed ingress", "Election or leader gate", "Fences"]
cwid = (CW - 4 * 0.3) / 5
for i, c in enumerate(chips):
    x = L + i * (cwid + 0.3)
    box(s, x, 5.5, cwid, 0.7, c, style="queue", tsize=14, kind=MSO_SHAPE.ROUNDED_RECTANGLE)
    if i < 4:
        arrow(s, [(x + cwid, 5.85), (x + cwid + 0.3, 5.85)], "a")
text(s, L, 6.3, CW, 0.65, [
    {"runs": [("Fence: ", {"bold": True}), ("an instance that can no longer trust its view of the log closes its edge or "
                                           "leader gate, releases its session, reports once and stops acting.", {})], "space": 2},
    {"runs": [("Topology: ", {"bold": True}), ("one document declares both kinds; clusterctl load-topology puts it on "
                                              "the log.", {})]}], size=13, color=A5)

s = content("Java client API", "One object, one call per duty cycle",
            "Both are from getting-started.md, trimmed to seqeron's part: connecting to the media driver is one Aeron "
            "call before them, and the idle strategy is Agrona's. The consumer is one object, ReplayerStreamReceiver, and one poll "
            "per duty cycle: it asks its member's Replayer for the history, delivers it, and switches to the live queue "
            "when the two meet, every frame once and in globalSeqNo order. The producer takes a façade, here "
            "Application: it holds the cluster session, follows the stream, confirms its own messages and resends what "
            "a failover lost, and opens its gate only on the leader's member, so the loop is the application's own logic. "
            "Published means placed: the message is sequenced when it comes back through onSequenced. onFenced means "
            "stop: the façade has already closed the gate and released the session, and the examples exit with 70, "
            "the repository's convention for fenced, restart me. Nothing here starts a thread: doWork() and poll() return a work count for the caller's idle "
            "strategy. C++ and C# have the same façades; C# takes an options class where Java takes a builder.")
CONSUMER = """ReplayerStreamReceiver receiver =
    new ReplayerStreamReceiver(
        23,                     // clientId
        event -> handle(event), // history, then live
        (leader, term, seqNo) -> {},
        () -> log("caught up"));
receiver.start(aeron, 0);       // member 0
while (running) {
    idle.idle(receiver.poll());
}"""
PRODUCER = """Application app = Application.builder()
    .sourceId(100).clientId(24).memberId(0)
    .egressChannel(EGRESS)
    .listener(listener).build();
app.start(aeron);
while (running) {
    final int work = app.doWork();
    if (pending && app.canPublish()) {
        // Declined: retry next cycle; Refused is final
        pending = app.publish(PAYLOAD_ID, msg, length)
            == Publish.Declined;
    }
    idle.idle(work);
}
// Listener: frames in order; a fence means stop
void onSequenced(Payload payload)
void onFenced(ClusterError fence, String detail)"""
for i, (t, sub, src) in enumerate([("Consumer", "the read path, no façade needed", CONSUMER),
                                   ("Producer", "a co-located application", PRODUCER)]):
    x = L + i * (5.9 + 0.33)
    text(s, x, 1.75, 5.9, 0.4, [{"runs": [(t, {"bold": True}), ("  " + sub, {"color": A5})]}], size=16)
    sh = shape(s, x, 2.2, 5.9, 3.4, "card", name="Code")
    tf = sh.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for side in ("left", "right", "top", "bottom"):
        setattr(tf, "margin_" + side, I(0.18))
    fill_tf(tf, [{"runs": [(ln.split("//")[0], {}), ("//" + ln.split("//", 1)[1], {"color": A1})]
                  if "//" in ln else [(ln, {})]} for ln in src.split("\n")], size=11)
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.name = "Menlo"
pts = [("One thread, yours. ", "doWork() returns a work count; dispatch allocates nothing once warm, tested in C#."),
       ("Not in your code: ", "the session, recovery, confirmation and resend, the leader gate, the fences."),
       ("Java, C++ and C#: ", "the same façades; C# takes an options class where Java takes a builder.")]
for i, (b, d) in enumerate(pts):
    x = L + i * (3.84 + 0.305)
    card(s, x, 5.85, 3.84, 0.95, [{"runs": [(b, {"bold": True}), (d, {})], "size": 13}], margin=0.15)


s = content("The gateway election", "Every input comes from the log; the output is a sequenced frame",
            "Walk the tape left to right. The topology rows arrive; the last one has remaining 0, which is the "
            "completeness edge, and the sequencer synthesizes GatewayActive for each rank-0 instance. GW-A announces "
            "itself with GatewayStarted, which binds its cluster session, and opens connections. When A's session "
            "closes, through a crash, network loss or shutdown, the sequencer designates B. B's GatewayStarted carries a "
            "first connection id above the highest it saw, and a consumer treats it as closing every connection still "
            "open under the pair, since a crashed instance never closes them. B rebuilt A's connections from the log while "
            "it was a hot standby.")
sc = CW / 980
f = Fig(s, L - 10 * sc, 1.75, sc)
f.label(10, 0, 400, 26, "log, in globalSeqNo order →", size=12, align=PP_ALIGN.LEFT)
tape_frames = [("GatewayRegistered", "GW-A · rank 0", "box"), ("GatewayRegistered", "GW-B · rank 1 · last", "box"),
               ("GatewayActive", "GW-A · bootstrap", "syn"), ("GatewayStarted", "GW-A · binds session", "box"),
               ("ConnectionOpened", "ids 1 … n, by GW-A", "box"), ("GatewayActive", "GW-B · A's session closed", "syn"),
               ("GatewayStarted", "GW-B · first id > n", "box")]
for k, (t, d, st) in enumerate(tape_frames):
    f.box(10 + k * 138, 34, 128, 82, t, d, st, tsize=11, ssize=10)
f.label(10, 140, 60, 30, "GW-A", size=13, color=TX, align=PP_ALIGN.LEFT, bold=True)
f.box(424, 142, 276, 26, "serving", style="bar", tsize=11, tcolor=BG)
f.label(706, 140, 200, 30, "crashed: session closes", size=11, align=PP_ALIGN.LEFT)
f.label(10, 190, 60, 30, "GW-B", size=13, color=TX, align=PP_ALIGN.LEFT, bold=True)
f.box(70, 192, 630, 26, None, style="dim")
f.label(80, 190, 600, 30, "hot standby: follows the queue, rebuilds A's connections from the log", size=11, color=TX,
        align=PP_ALIGN.LEFT)
f.box(838, 192, 128, 26, "serving", style="bar", tsize=11, tcolor=BG)
y0 = 1.75 + 240 * sc
shape(s, L, y0 + 0.1, 0.3, 0.3, "syn", name="Legend")
text(s, L + 0.4, y0 + 0.05, 5, 0.4, "synthesized by the sequencer", size=12, color=A5, anchor=M)
shape(s, L + 4.2, y0 + 0.1, 0.3, 0.3, "box", name="Legend")
text(s, L + 4.6, y0 + 0.05, 6, 0.4, "submitted by clusterctl or an instance", size=12, color=A5, anchor=M)
card(s, L, y0 + 0.7, CW, 0.95, [{"runs": [("Every gateway instance and every application sees the same decision at the same globalSeqNo. ", {"bold": True}),
                                          ("No two instances can act as active on inconsistent information.", {})], "size": 15}],
     margin=0.22, anchor=M)

s = content("Leader-only work, through the log", "Outstanding work is replicated; dispatch is local",
            "Every application instance records a request as outstanding from its sequenced request until its sequenced reply. Only "
            "the leader's instance dispatches; it replies under the requester's sourceId and connectionId so the "
            "originating gateway can route it. The request is discharged when the reply is sequenced, not when it is "
            "sent. Every LeadershipChanged closes the gate, including one naming the same member, because a reply sent "
            "during the election may have been lost with the old leader's uncommitted log. A newly opened gate "
            "dispatches every outstanding request in globalSeqNo order. Delivery is at least once across a failover; a "
            "reply that is a pure function of its request is byte-identical on re-emission and deduplicated by key.")
sc = CW / 960
f = Fig(s, L, 1.7, sc)
f.host(10, 120, 120, 50, "External")
f.box(170, 120, 150, 50, "Gateway", "active instance")
f.arrow([(130, 138), (170, 138)])
f.arrow([(170, 154), (130, 154)])
f.box(360, 120, 110, 50, "seq N", "request", "queue")
f.arrow([(320, 145), (360, 145)])
f.box(510, 40, 180, 50, "Application · member 0", "outstanding: N", tsize=11)
f.box(510, 120, 180, 50, "Application · member 1", "leader: reply(requester, conn)", "lead", tsize=11)
f.box(510, 200, 180, 50, "Application · member 2", "outstanding: N", tsize=11)
for ty in (65, 145, 225):
    f.arrow([(470, 145), (510, ty)], "d")
f.box(730, 120, 110, 50, "seq M", "reply", "queue")
f.arrow([(690, 145), (730, 145)], "a")
f.label(845, 112, 115, 66, "every instance: N discharged", size=11, align=PP_ALIGN.LEFT)
f.arrow([(785, 170), (785, 278), (245, 278), (245, 170)], "a")
f.label(300, 280, 440, 24, "gateway keeps its own sourceId, routes by connectionId", size=11)
pts3 = [("Replicated", "Outstanding from sequenced request to sequenced reply, on every instance."),
        ("Gate closes", "On every LeadershipChanged, even for the same leader, and at a fence."),
        ("At least once", "Deterministic replies make duplicates identical; drop by key.")]
for i, (t, d) in enumerate(pts3):
    x = L + i * (3.84 + 0.305)
    card(s, x, 5.7, 3.84, 1.05, [{"runs": [(t + "  ", {"bold": True}), (d, {})], "size": 13}], margin=0.15, anchor=M)

s = content("Confirmed ingress", "A failover can lose uncommitted ingress silently; the producer's own queue finds it",
            "A send returns once the frame is on the leader's ingress publication; the cluster confirms nothing. Frames "
            "the old leader had not committed are lost, and the session survives, so neither side sees a fault: about "
            "5,500 frames at 100 µs pacing in the failover harness. PendingSends holds each placed frame until the "
            "producer's own queue shows it, matched by cluster session. A LeadershipChanged with a newer term identifies "
            "exactly which were lost: the newest frames stamped with the old term that have not appeared. Nothing new is "
            "placed until they are resent, oldest first. Each frame is then sequenced exactly once, in order. The limit: "
            "this is in memory, so a restarted producer or a promoted standby starts with nothing pending, and a replaced "
            "session or a superseded gateway instance drops what was pending.")
text(s, L, 2.2, 1.8, 0.6, "Producer places", size=14, bold=True, anchor=M)
text(s, L, 3.6, 1.8, 0.6, "Own queue shows", size=14, bold=True, anchor=M)
cw_, cg = 0.75, 0.08
X0 = 2.5
for k in range(6):
    x = X0 + k * (cw_ + cg)
    lost = k >= 3
    box(s, x, 2.2, cw_, 0.6, f"f{k + 1}", style="miss" if lost else "queue", tsize=13, tcolor=A2 if lost else TX)
for k in range(3):
    box(s, X0 + k * (cw_ + cg), 3.6, cw_, 0.6, f"f{k + 1}", style="queue", tsize=13)
box(s, X0 + 3 * (cw_ + cg), 3.6, 2 * cw_ + cg + 0.6, 0.6, "LeadershipChanged", "term T+1", style="syn", tsize=12, ssize=10)
R0 = 8.0
for k in range(4):
    x = R0 + k * (cw_ + cg)
    lab = f"f{k + 4}"
    box(s, x, 2.2, cw_, 0.6, lab, style="syn" if k < 3 else "queue", tsize=13)
    box(s, x, 3.6, cw_, 0.6, lab, style="queue", tsize=13)
arrow(s, [(X0 + 3 * (cw_ + cg), 2.92), (X0 + 3 * (cw_ + cg), 3.02), (X0 + 6 * (cw_ + cg) - cg, 3.02),
          (X0 + 6 * (cw_ + cg) - cg, 2.92)], "w", head=False)
text(s, X0 + 3 * (cw_ + cg), 3.05, 2.6, 0.35, "lost with the old leader", size=11, color=A2)
text(s, R0, 1.75, 4.5, 0.4, "resent oldest first, then new work", size=12, color=A2, bold=True)
text(s, 7.42, 2.2, 0.55, 0.6, "hold", size=12, color=A5, align=C, anchor=M)
for i, (t, d) in enumerate([("Placed is not sequenced", "Each frame is held until the producer's own queue shows it."),
                            ("A term change names the losses", "The newest frames of the old term not yet seen."),
                            ("Limit", "In memory: a restarted producer or promoted standby starts empty; a replaced session drops what was pending.")]):
    x = L + i * (3.84 + 0.305)
    card(s, x, 4.75, 3.84, 1.85, [{"text": t, "bold": True, "size": 15, "space": 6, "color": A2 if i == 2 else TX},
                                  {"text": d, "size": 14}], margin=0.2)

s = content("Application snapshots", "Local files the log confirms; the cluster itself takes none",
            "The sequencer sequences SnapshotStarted as a round's cut, on an operator request or the policy interval. "
            "Every instance of a participating source serializes its state at that cut into its own local file. The "
            "instance that may publish submits SnapshotEnd with record count, length and CRC32C. Every instance compares "
            "its own file with the sequenced end and fences on a difference, which also makes each round a divergence "
            "check across instances; a match confirms the file, which replaces every older one. A restarting instance "
            "restores its confirmed file once a replay from the file's cut position opens on that cut, and replays only "
            "what follows; with no such file it replays from 1. The Replayer holds no snapshot state. "
            "This shortens an application's restart, not a member's.")
flow = [("Begin", "Sequencer sequences SnapshotStarted"), ("Serialize", "Every instance writes its own local file"),
        ("End", "The publishing instance submits SnapshotEnd: count, length, CRC32C"),
        ("Compare", "Every instance checks its file: a match confirms it, a difference fences"),
        ("Restore", "Newest confirmed file, then only what follows the cut")]
fw = (CW - 4 * 0.35) / 5
for i, (t, d) in enumerate(flow):
    x = L + i * (fw + 0.35)
    shape(s, x, 2.0, fw, 2.9, "syn" if i == 0 else "card", name="Step")
    num_circle(s, x + 0.2, 2.2, i + 1, d=0.45, size=14, style="warm" if i == 0 else "accent")
    text(s, x + 0.2, 2.8, fw - 0.4, 0.45, t, size=18, bold=True, head=True)
    text(s, x + 0.2, 3.3, fw - 0.4, 1.5, d, size=14)
    if i < 4:
        arrow(s, [(x + fw, 3.45), (x + fw + 0.35, 3.45)], "a")
card(s, L, 5.3, CW, 1.3, [{"text": "Application snapshots bound a client's restart to the messages since the last one "
                                   "and are the intended boundary for application version upgrades.", "size": 16}],
     margin=0.25, anchor=M)

# ================================================================= FAILURES
section("Failures and verification")
s = content("Failure model", "Each failure and its recovery",
            "Walk the table top to bottom. Note what is absent: there is no recovery procedure that is not a replay. "
            "globalSeqNo, the gateway list and the active instance are all functions of the log. "
            "The timings are the defaults. A follower starts an election after 200 ms without a leader heartbeat; how "
            "long the election takes is not measured. A member exits after 1 s of back-pressure with no recording "
            "progress, and a stopped recording is caught by the 1 Hz heartbeat's liveness check. A slow consumer's "
            "driver drops it within about 0.2 s. A missed frame shows at the next frame, which the heartbeat bounds at "
            "about a second. The cluster closes a silent client's session after the 1 s session timeout, which is what "
            "promotes a crashed gateway's standby. A fence on a mismatch is immediate; the stall fences wait 20 s on the "
            "queue and 60 s in recovery, and only on an instance that acts. A fenced gateway's standby takes over at "
            "once, but nothing replaces a fenced leader's application replica: that application's leader-only work "
            "waits until the replica restarts and catches up, or leadership moves. A gateway host's relay moves after 3 s, three "
            "heartbeats, with no frame.")
table(s, L, 1.8, [2.75, 2.95, 2.68, 3.75], [
    ["Failure", "What happens", "Detected in (defaults)", "Recovery"],
    ["Leader lost", "Election; LeadershipChanged", "200 ms heartbeat timeout", "Next globalSeqNo; producers resend"],
    ["Member lost", "Others keep quorum", "Not needed; no pause", "Restart replays the full log"],
    ["Archive stops", "Member exits with code 70", "1 s without progress", "As for a lost member"],
    ["Slow consumer", "Dropped; queue unaffected", "About 0.2 s", "Resumes over the gap"],
    ["Frame missed on queue", "Next frame reveals the gap", "Next frame, about 1 s", "Resume, holding live frames"],
    ["Gateway instance fails", "Its session closes", "1 s session timeout", "Standby designated"],
    ["Client loses its session", "The cluster closes it", "1 s session timeout", "New session; announced gateway fences"],
    ["Gateway instance fences", "Edge closed, session released", "At once; stalls 20 s, 60 s", "Standby takes over at once"],
    ["Leader's application fences", "Gate closed, session released", "At once; stalls 20 s, 60 s", "Its restart, or a new leader's replica"],
    ["Relay's member lost", "The relay hears nothing", "3 s", "Next member, same position"],
    ["Majority lost", "Nothing commits; queue silent", "Acting instances at 20 s", "Resumes with a majority"],
], size=13, rowh=0.41)

s = content("Verification", "A harness per failure, against a live cluster",
            "Every failure in the table has an end-to-end harness against a live cluster. chaos-runner injects "
            "randomized faults from a printed seed and ends by checking every member's recording is gap-free and at the "
            "same high mark; every member kill also kills the media driver its clients share, and they must exit. "
            "Network faults are injected with eBPF against the containerized cluster. The harnesses drive the Java client, "
            "and the C# client in two of its own. Beneath them, the Java, C++ and C# unit suites test the same decisions "
            "case for case, and their property tests draw the same fault sequence from a seed. The C++ client is "
            "covered by its unit suite alone: no harness runs it against a live cluster. Killing a standalone aeronmd "
            "is not exercised either.")
harn = [("chaos-runner", "random faults from a seed; recordings gap-free and equal"),
        ("failover soak", "repeated leader kills in containers"),
        ("network faults", "eBPF: partitions, lost repair, slow links, dropped ballots"),
        ("confirmed ingress", "every frame once, in order, across a kill"),
        ("gateway host", "relay moves when its member dies"),
        ("session replace", "paused client, lost quorum, paused gateway pair"),
        ("snapshots", "restore, handover, leader kill"),
        ("C# client", "Linux and Windows"),
        ("unit suites", "Java, C++ and C#, case for case; seeded property tests")]
for i, (t, d) in enumerate(harn):
    y = 1.85 + i * 0.47
    text(s, L, y, 3.0, 0.5, t, size=16, bold=True, anchor=M)
    text(s, L + 3.1, y, CW - 3.1, 0.5, d, size=16, color=A5, anchor=M)
text(s, L, 6.2, CW, 0.45, [{"runs": [("Not exercised: ", {"bold": True, "color": A2}),
                                    ("the C++ client against a live cluster; killing a standalone aeronmd.", {})]}],
     size=16)

s = content("Monitoring", "Each member's counters, scraped by Prometheus, on one Grafana dashboard",
            "One exporter runs beside each member, sharing its Aeron directory, and serves the member's counters over "
            "HTTP in Prometheus text format, read live from the CnC file; it needs no cluster session. Prometheus "
            "pulls every exporter from a static target list, and its own up metric is the member's reachability. "
            "Grafana is provisioned from files with one dashboard: a panel per metric family plus node apply lag, now "
            "minus the last ClusterHeartbeat, which is the one place a member publishing a contiguous but stale queue "
            "shows up. Clients publish counters too: recovery stalled, a gateway instance's open connections, an "
            "instance's confirmed snapshot round, and an application's own in type ids 5201 to 5297, labelled by member "
            "and client. Not provided: alert rules, service discovery, authentication on "
            "the metrics endpoint.")
pipe = [("Member", "Aeron counters in cnc.dat", "box"), ("metrics-exporter", "HTTP · port 9400 + member id", "syn"),
        ("Prometheus", "pulls every exporter", "box"), ("Grafana", "one provisioned dashboard", "box")]
pw = 2.6
pg = (CW - 4 * pw) / 3
for i, (t, d, st) in enumerate(pipe):
    x = L + i * (pw + pg)
    box(s, x, 1.85, pw, 0.9, t, d, style=st, tsize=15, ssize=11)
    if i < 3:
        arrow(s, [(x + pw, 2.3), (x + pw + pg, 2.3)])
groups = [("Sequencer", ["globalSeqNo and current leader", "Apply lag: now − last heartbeat",
                         "Rejected ingress, leadership changes", "Queue back-pressure and stall", "Gateway promotions"]),
          ("Replayer", ["Ready, archive errors", "Active and pending replays",
                        "Throttled: the recorder is behind"]),
          ("Clients", ["Recovery stalled, per client", "Open connections, per gateway instance",
                       "Confirmed snapshot round, per instance",
                       "An application's own counters, type ids 5201–5297",
                       "Labelled member and client"])]
for i, (t, pts) in enumerate(groups):
    x = L + i * (3.84 + 0.305)
    card(s, x, 3.15, 3.84, 2.8, [{"text": t, "size": 18, "bold": True, "head": True, "space": 8}]
         + [{"text": p, "size": 14, "bullet": True, "space": 4} for p in pts], margin=0.22)
text(s, L, 6.2, CW, 0.4, [{"runs": [("Not provided: ", {"bold": True, "color": A2}),
                                    ("alert rules, service discovery, authentication on the metrics endpoint.", {})]}], size=14)

# ================================================================= EXPERIENCE
section("Building on Aeron")
s = content("What Aeron does not tell you", "Behavior learned the hard way, and what seqeron does about it",
            "Each row is a behavior of Aeron that is correct by its own contract and easy to build on wrongly, and "
            "each is already on a slide. Confirmed ingress answers the first; the queue's bounded, reliable offer the "
            "fourth and fifth; the replay sequences the seventh. A cluster callback that throws has had its log "
            "position advanced by Image.boundedControlledPoll and AgentRunner keeps the agent alive, so the frame is "
            "dropped and nothing stops; a member that cannot continue exits with code 70 instead. The client side is "
            "the same: Image.poll gives a handler's exception to the error handler and advances, so a client records "
            "the fault and raises it from its own duty cycle. CLOSED from an ingress offer is an election in progress, "
            "not a dead session. Aeron.NET's defaults are Aeron's, whose new-leader timeout is shorter than an "
            "election, which csharp-client-test.sh catches on every pull request.")
table(s, L, 1.75, [3.7, 4.3, 4.13], [
    ["Aeron", "Consequence", "seqeron"],
    ["An accepted offer is not a committed one", "A failover loses it; the session survives, so nobody sees it",
     "Confirmed on the producer's own queue"],
    ["A cluster callback that throws", "Its position has advanced: the frame is dropped, the agent runs on",
     "Never throws; exits with code 70"],
    ["A handler that throws in Image.poll", "The error handler gets it; the position advances anyway",
     "Recorded, raised from doWork()"],
    ["Back-pressure is the caller's", "A service thread that blocks stalls the log and the cluster clock",
     "A bounded retry watching RecordingPos"],
    ["A stopped recording", "Back-pressures nothing: untethered readers keep it connected",
     "The 1 Hz heartbeat's liveness check"],
    ["CLOSED from an ingress offer", "Looks terminal; it is an election in progress", "Keeps spinning: IngressStallPolicy"],
    ["A bounded replay of a live recording", "Does not close at its bound", "Completion by position"],
    ["Aeron.NET's defaults", "A new-leader timeout shorter than an election", "Every timeout set explicitly"],
    ["No C++ cluster client", "Nothing to build on", "Written against the wire protocol"],
], size=12, rowh=0.48)


s = prs.slides.add_slide(L_SECTION)
_track(s)
tape(s, 0.8, 2.15, 6, 0)
for ph in s.placeholders:
    tf = ph.text_frame
    tf.margin_left = tf.margin_right = 0
    if ph.placeholder_format.idx == 0:
        tf.vertical_anchor = MSO_ANCHOR.TOP
        fill_tf(tf, [{"text": "Discussion", "size": 44, "bold": True, "color": BG, "head": True}])
    else:
        fill_tf(tf, [{"text": "Further reading in doc/", "size": 20, "color": A3}])
text(s, 0.8, 5.35, 11.7, 1.6, [
    {"runs": [("overview.md", {"bold": True, "color": BG}), ("  the design   ", {}),
              ("seqeron-protocol-spec.md", {"bold": True, "color": BG}), ("  normative rules", {})], "space": 6},
    {"runs": [("fault-tolerance.md", {"bold": True, "color": BG}), ("  each failure   ", {}),
              ("client-api.md", {"bold": True, "color": BG}), ("  the client tier   ", {}),
              ("snapshot.md", {"bold": True, "color": BG}), ("  application snapshots", {})], "space": 6},
    {"runs": [("ops.md", {"bold": True, "color": BG}), ("  ports, counters, runbooks   ", {}),
              ("clusterctl.md", {"bold": True, "color": BG}), ("  the operator tool", {})]}], size=16, color=A3)
notes(s, "The last slide presented; the appendix is backup, for questions or a meeting that runs long. Open the "
         "floor. The decisions slide, after failover at the start of the appendix, is the natural place to return to.")

# ================================================================= APPENDIX
section("Appendix")
s = content("Failover", "Where a client switches when its primary cannot be reached; the cluster's own election aside",
            "The cluster tier's own failover, the Raft election and a member that exits rather than lose its "
            "recording, is on the failure model slide. These are the client tier's. Ingress: a co-located producer "
            "uses IPC to its own member and UDP to the leader otherwise; Java and C# replace the session when "
            "leadership moves off the co-located member, C++ swaps only its publication and follows a follower's "
            "redirect. A send spins through an election rather than dropping the frame. A session the cluster closed "
            "is replaced, one attempt a second, and the instance fences after 20 s without one; a gateway instance "
            "that announced itself on the lost session fences at once, since its sibling may already be designated. "
            "PendingSends resends what the old leader lost, exactly once and in order. The gateway pair and the "
            "leader-only work of applications move with the log. A gateway host's relay moves to the next member's "
            "archive and resumes at the same position, valid because every member's active recording starts at "
            "globalSeqNo 1. A client's own Replayer, queue and media driver have no alternative: an instance that "
            "acts fences. A fenced gateway instance's standby takes over at once; nothing replaces a fenced "
            "leader's application replica, so that application's leader-only work waits until the replica "
            "restarts and catches up, or leadership moves to another member.")
table(s, L, 1.8, [3.1, 3.7, 5.33], [
    ["What", "Fails over to", "Trigger"],
    ["Ingress path", "UDP to the leader", "Own member not leading, or the IPC connect timed out"],
    ["A send during an election", "The new leader's ingress", "NewLeader; the send spins rather than drop the frame"],
    ["Cluster session", "A new session", "The cluster closed it; one attempt a second, a fence after 20 s"],
    ["Ingress the old leader lost", "A resend to the new leader", "A newer term's LeadershipChanged"],
    ["Gateway pair", "The next instance by rank", "Session closed, 5 s without GatewayStarted, or clusterctl activate"],
    ["Leader-only work", "The new leader's replica", "LeadershipChanged opens its gate; outstanding work redispatched"],
    ["Gateway host's queue relay", "The next member's archive", "3 s without a frame, or a failed read; same position"],
], size=13, rowh=0.55)
card(s, L, 6.28, CW, 0.68, [{"runs": [("No failover: ", {"bold": True, "color": A2}),
                                      ("a client's own Replayer, queue and media driver. An instance that acts fences: "
                                       "a gateway's standby takes over at once, but a leader's application waits for "
                                       "its own restart or a new leader.", {})]}],
     style="plain", margin=0.1, anchor=M)

s = content("Decisions and what they cost", "The choices worth challenging",
            "Lead the discussion here. The most consequential is no cluster snapshots: it buys complete local history "
            "on every member and costs member recovery time and archive size that grow with uptime; the 1 Hz heartbeat "
            "alone is about 86,400 frames a day. The single stream is the other: one order for everything with no "
            "coordination, at the price of one sequencer's throughput for all producers.")
table(s, L, 1.8, [3.75, 4.25, 4.13], [
    ["Decision", "Buys", "Costs"],
    ["No cluster snapshots", "Complete history on every member", "Recovery and archive grow with uptime"],
    ["Each member records its queue", "Local recovery for every client", "Each member stores everything"],
    ["Untethered consumers", "Nobody slows the sequencer", "A slow consumer pays a replay"],
    ["Stop on recording failure", "Never serves history it lacks", "Costs a member, not a frame"],
    ["Opaque payloads", "Applications version alone", "No content routing or checks"],
    ["One ordered stream", "No coordination anywhere", "One sequencer's throughput"],
    ["Majority, no fsync", "Memory and network latency", "Correlated power loss"],
], size=14, rowh=0.56)

s = content("Limits today", "Stated plainly, so nobody discovers them in production",
            "Seven members is bounded by the 70-port cluster block. Payloads above 8,884 bytes are split by the "
            "application, which keeps a frame within one IPC MTU. Four concurrent replays per member, one while the "
            "recording trails the queue; more wait. "
            "Nothing on the wire is authenticated or encrypted: the network is the access control, so the port block "
            "belongs inside a trusted network. No latency, throughput or sizing figures have been measured on production "
            "hardware yet. A change to seqeron's own protocol moves every member, gateway host and client together on "
            "purged archives; applications upgrade independently.")
lims = [("7", "members at most", "bounded by the 70-port cluster block"),
        ("8,884 B", "payload", "larger payloads are split by the application"),
        ("4", "replays per member concurrently", "1 while the recording is behind; more wait"),
        ("None", "wire security", "no authentication or encryption on the wire"),
        ("Unmeasured", "performance", "no latency or throughput figures yet"),
        ("All at once", "framing protocol upgrade", "seqeron's own protocol, on purged archives")]
for i, (big, lab, d) in enumerate(lims):
    col, row = i % 3, i // 3
    x, y = L + col * (3.84 + 0.305), 1.85 + row * (2.3 + 0.25)
    card(s, x, y, 3.84, 2.3, [{"text": big, "size": 24, "bold": True, "color": A1, "head": True, "space": 4},
                              {"text": lab, "size": 15, "bold": True, "space": 4},
                              {"text": d, "size": 14, "color": A5}], margin=0.22)

s = content("Non-functional requirements", "What seqeron gives today, attribute by attribute",
            "Each row is what the design and the harnesses establish. Latency and throughput have never been measured "
            "on production hardware; a deployment benchmarks on its own, at its peak rather than its daily volume. A "
            "full day of Nasdaq's US ITCH feed peaks at 988,711 messages in the closing second, about 65 times its "
            "15,162 a second over regular hours; at the same ratio, the Nordic cash-equity market's estimated average "
            "of 800 a second (the sizing slide) peaks near 50,000 a second. That is an estimate, not a requirement "
            "anyone has set. Failover starts after 200 ms without a leader heartbeat by default "
            "(sequencer.leaderHeartbeatTimeoutMs); how long the election and the producers' resend take is not "
            "measured. Data loss is nil for any minority's crash, and for every member process crashing, because the "
            "page cache survives a process; a power loss across a majority of hosts before write-back can lose "
            "committed messages, since nothing is fsynced.")
table(s, L, 1.8, [3.0, 9.13], [
    ["Attribute", "Today"],
    ["Latency", "Ingress to the producer's own queue: not measured"],
    ["Throughput", "One stream: not measured; the target is the peak, an estimated 50,000/s at a Nordic close"],
    ["Fault tolerance", "A minority of members: 1 of 3, 2 of 5, 3 of 7"],
    ["Failover", "Election after 200 ms without a leader heartbeat; duration not measured"],
    ["Data loss", "None for a minority's crash; a majority's power loss can lose committed messages"],
    ["Member restart", "Full-log replay: grows with uptime"],
    ["Client restart", "Newest confirmed snapshot, then the tail; else the full log"],
    ["Retention", "Everything, on every member; truncation designed, not built"],
], size=14, rowh=0.52)

s = content("Sizing the archive", "Every member records the Raft log and the queue, and keeps both",
            "Per message, every member writes the Raft log entry and the queue frame, each in a 32-byte-aligned Aeron "
            "frame with a 32-byte header. The log entry carries Aeron Cluster's 32-byte session header and seqeron's "
            "28-byte ingress envelope; the queue frame carries the 44-byte sequenced envelope. The illustration is an "
            "estimate of Nasdaq's Nordic cash-equity market on INET, and arithmetic, not a measurement. Nasdaq's May "
            "2024 statistics report 611,613 share trades a trading day across the Nordic and Baltic exchanges. Order "
            "messages per trade (entries, amends, cancels) were measured at 11 to 16 on the 50 most-traded Nordic "
            "stocks in 2010 to 2012; assuming 40 today gives about 25 million messages a day, some 800 a second "
            "averaged over the trading day, with the open and the close far above that. The payload is 32 bytes: a "
            "full day of Nasdaq's US ITCH 5.0 feed averages 28.5 bytes a message, and Nordic OUCH 5's order entry "
            "messages, weighted by that day's mix, 28 to 37 (Enter Order 41 bytes before optional fields, Replace 25, "
            "Cancel 15). That is 128 plus 128 bytes a message: 6.4 GB a day on each member, 1.6 TB over 250 trading "
            "days. At these sizes the 32-byte alignment dominates: a 64-byte payload would make it 8 GB. "
            "Derivatives, ETFs and warrants are not included. Every start of a member replays the log from globalSeqNo 1 into "
            "a new queue recording, and the earlier recordings stay on disk unread, so restarts add a full copy each. "
            "A member's restart time is its Raft log divided by the replay rate, which is not measured. Cluster "
            "snapshots and truncation (snapshot.md, section 10) would bound both, and are designed, not built.")
card(s, L, 1.8, 5.9, 2.5, [
    {"text": "Per message, on every member", "size": 18, "bold": True, "head": True, "space": 10},
    {"runs": [("Raft log entry   ", {"bold": True}), ("align32(92 + payload)", {"color": A1, "bold": True})],
     "size": 16, "space": 6},
    {"runs": [("Queue frame   ", {"bold": True}), ("align32(76 + payload)", {"color": A1, "bold": True})],
     "size": 16, "space": 10},
    {"text": "Aeron frame header 32 B · cluster session header 32 B · seqeron envelope 28 B in, 44 B out",
     "size": 13, "color": A5}], margin=0.25)
card(s, L + 6.23, 1.8, 5.9, 2.5, [
    {"text": "Illustration: INET Nordic equities", "size": 18, "bold": True, "head": True, "space": 10},
    {"runs": [("25 M", {"size": 26, "bold": True, "color": A1, "head": True}),
              ("  messages a day at 32 B: 256 B each", {"size": 14})], "space": 4},
    {"runs": [("6.4 GB", {"size": 26, "bold": True, "color": A1, "head": True}),
              ("  a day on each member; 1.6 TB a year", {"size": 14})], "space": 4},
    {"text": "600k trades a day × 40 messages; 32 B from ITCH and OUCH", "size": 13, "color": A5}], margin=0.25)
for i, paras in enumerate([
        [{"text": "Each restart adds a copy", "size": 15, "bold": True, "space": 4},
         {"text": "A member re-records the queue from globalSeqNo 1; earlier recordings stay on disk, unread.",
          "size": 13}],
        [{"text": "Restart time", "size": 15, "bold": True, "space": 4},
         {"text": "Raft log ÷ replay rate; the replay rate is not yet measured.", "size": 13}],
        [{"text": "Bounded later", "size": 15, "bold": True, "space": 4},
         {"text": "Cluster snapshots and truncation: designed (snapshot.md §10), not built.", "size": 13}]]):
    x = L + i * (3.84 + 0.305)
    card(s, x, 4.55, 3.84, 1.6, paras, style="plain", margin=0.18)

s = content("Deployment and disaster recovery", "What a host needs, where members go, and what is not yet designed",
            "A member host runs one JVM and needs local disk for the archive and the Raft log, and on Linux a tmpfs "
            "Aeron directory. Placement follows from Raft: a member per failure domain, so no single rack, power feed or "
            "zone holds a majority. That is also the mitigation for the no-fsync durability model, which loses committed "
            "messages only when a majority of hosts lose power together. Spreading members further apart adds the "
            "round trip between them to every message, since a commit waits for a majority. Archives can be copied off "
            "site with Aeron Cluster Backup, which nothing in seqeron exercises yet. The rest of disaster recovery is not designed: whether a second site is a "
            "stretched cluster or a standby, how a restore is verified, and the RTO and RPO for losing a site, are all "
            "open.")
dep = [("A member host", TX, [[("One JVM (JDK 21): consensus, sequencer, driver, archive, Replayer", {})],
                              [("Archive and Raft log on local disk; Aeron directory on tmpfs", {})],
                              [("A 70-port block from 9300; metrics on 9400 + member id", {})],
                              [("Clients beside a member, or on a gateway host that relays its queue", {})]]),
       ("Placement", TX, [[("Each member in its own failure domain: rack, power, zone", {})],
                          [("A commit waits for a majority: the round trip between members is on every message", {})],
                          [("Separate power keeps committed messages safe without fsync", {})]]),
       ("Disaster recovery", A2, [[("Archives copied off site: Aeron Cluster Backup, not yet exercised", {})],
                                  [("Open: ", {"bold": True}), ("a second site, stretched cluster or standby", {})],
                                  [("Open: ", {"bold": True}), ("how a restore is verified", {})],
                                  [("Open: ", {"bold": True}), ("recovery time and data loss for a site loss", {})]])]
for i, (t, c, pts) in enumerate(dep):
    x = L + i * (3.84 + 0.305)
    card(s, x, 1.85, 3.84, 3.6, [{"text": t, "size": 18, "bold": True, "head": True, "color": c, "space": 10}]
         + [{"runs": p, "size": 14, "bullet": True, "space": 8} for p in pts], margin=0.22)

s = content("Security posture", "Crash faults, not malicious ones: the network is the access control",
            "The threat model is crash faults. The cluster checks that a message is well formed, not who sent it, and "
            "authentication belongs at the system's external edges. The one identity check is the gateway rule S-6: "
            "once designated, an instance is held to the cluster session it announced itself on. Everything else "
            "follows from the network: whoever reaches the port block can submit, including the operator's events, read "
            "and delete recordings, and send Raft traffic. Payloads are stored unencrypted on every member and never "
            "deleted short of a purge, so personal data in a payload cannot be erased individually.")
table(s, L, 1.8, [2.5, 5.1, 4.53], [
    ["Surface", "Today", "Required of the deployment"],
    ["Ingress", "Anyone who reaches it can submit, operator events included", "Only producer hosts reach the ports"],
    ["Producer identity", "Checked for form, not origin; a gateway is held to its session", "Authenticate at the external edges"],
    ["Archive control", "Recordings can be read and deleted", "Member ports inside the boundary"],
    ["On the wire", "No encryption", "An encrypted, authenticated layer beneath"],
    ["At rest", "Unencrypted payloads on every member, kept indefinitely", "Disk encryption; classify payloads"],
    ["Erasure", "Not possible short of a purge", "No personal data in payloads"],
    ["Metrics", "No authentication", "Inside the boundary"],
], size=13, rowh=0.52)

s = content("Risks and mitigations", "Each risk, and what reduces it",
            "The first three matter most for a production deployment: "
            "unmeasured performance, an archive that grows without bound, and history lost on a framing protocol "
            "upgrade. A site loss would be covered by copying the archives off site with Aeron Cluster Backup, which "
            "nothing in seqeron exercises yet. "
            "Several clusters can scale out throughput, but nothing orders messages across them.")
table(s, L, 1.8, [4.0, 8.13], [
    ["Risk", "Mitigation"],
    ["Performance falls short", "Benchmark the closing peak on the target hardware"],
    ["Disk fills; restarts grow long", "Size by the formula; alert on disk; cluster snapshots (designed)"],
    ["Framing protocol upgrade purges history", "Export history to a system of record before each upgrade"],
    ["Majority power loss", "Members on separate power"],
    ["Unauthenticated ports", "Segmented network; an encrypted layer beneath"],
    ["Site loss", "Archives copied off site: Aeron Cluster Backup, not yet exercised"],
    ["One stream's ceiling", "Several clusters scale out"],
], size=14, rowh=0.52)

s = content("Glossary", "The terms this deck uses, one line each",
            "Reference for the whole deck. Rule ids cite seqeron-protocol-spec.md, whose letters are F frame, "
            "T transport, S sequencer, P payload, C registration, R replay, V versioning, E encoding and A application.")
for x, rows in ((L, [
    ["Application", "Deterministic service: one per member, or remote"],
    ["Cluster clock", "ClusterHeartbeat: 1 Hz, on consensus time"],
    ["Confirmed ingress", "Held until the producer's own queue shows it"],
    ["Façade", "A client's one entry point, Gateway or Application"],
    ["Fence", "Stops acting: its view of the log is not trusted"],
    ["Gate", "Lets an application instance publish; shut each term"],
    ["Gateway", "Edge producer pair; the cluster designates one"],
    ["Gateway host", "Clients, no member; relays a member's queue"],
    ["globalSeqNo", "The sequence number: from 1, no gaps, no reuse"],
    ["Member", "One cluster node: one JVM, its own archive"],
]), (L + 6.23, [
    ["payloadId", "A payload's type; the cluster never decodes it"],
    ["Producer", "Submits on ingress: a gateway or an application"],
    ["Queue", "A member's stream of sequenced frames: IPC 205"],
    ["Recording", "The archive's copy of a member's queue"],
    ["Replayer", "Serves history from the recording, per host"],
    ["Snapshot round", "An application snapshot the log cuts and confirms"],
    ["sourceId", "A producer's id: per application or gateway pair"],
    ["Term", "A Raft leadership term; one LeadershipChanged each"],
    ["Tethered", "Holds the publisher back; untethered is dropped"],
    ["S-3, A-4, V-1", "Spec rule ids: sequencer, application, versioning"],
])):
    table(s, x, 1.8, [1.75, 4.15], [["Term", "Meaning"]] + rows, size=12, rowh=0.44)

s = content("Aeron's transport model", "Publications, subscriptions, and one byte position that everything shares",
            "A publication writes into a log buffer: a memory-mapped file of three rotating terms. Offering is a copy "
            "into the current term plus an ordered write, with no lock and no system call. Position is the byte offset "
            "into the stream; flow control, loss detection, recording and replay all speak in positions. When the "
            "slowest subscriber the publication must wait for is a full window behind, offer returns back-pressured "
            "and the publisher decides what to do. Tethered versus untethered matters to seqeron: its archive is the "
            "queue's one tethered subscriber; every client is untethered.")
text(s, L, 1.9, 5.6, 4.6, [
    {"runs": [("Channel  ", {"bold": True}), ("aeron:udp (unicast, multicast) or aeron:ipc (shared memory)", {})], "bullet": True, "space": 10},
    {"runs": [("Stream id  ", {"bold": True}), ("names one stream on a channel", {})], "bullet": True, "space": 10},
    {"runs": [("Position  ", {"bold": True}), ("byte offset into the stream; the coordinate for flow control, loss, recording, replay", {})], "bullet": True, "space": 10},
    {"runs": [("Offer  ", {"bold": True}), ("a copy into a mapped term: no lock, no system call", {})], "bullet": True, "space": 10},
    {"runs": [("Back-pressure  ", {"bold": True}), ("returned to the caller; nothing queues without bound", {})], "bullet": True},
], size=15)
for i in range(3):
    x = 6.75 + i * 2.0
    box(s, x, 2.35, 1.85, 1.1, None, None, style="box")
    text(s, x, 2.4, 1.85, 0.3, f"term {i}", size=11, color=A5, align=C)
shape(s, 6.75, 2.72, 1.85, 0.73, "queue", name="Written")
shape(s, 8.75, 2.72, 1.3, 0.73, "queue", name="Written")
arrow(s, [(10.05, 1.95), (10.05, 3.55)], "a", head=False)
text(s, 8.75, 1.62, 2.6, 0.32, "publisher position", size=11, color=A1, align=C, bold=True)
arrow(s, [(7.15, 2.0), (7.15, 3.55)], "w", head=False)
text(s, 6.75, 3.6, 2.6, 0.32, "subscriber position", size=11, color=A2, bold=True)
arrow(s, [(7.15, 4.03), (10.05, 4.03)], "d", both=True)
text(s, 7.15, 4.05, 2.9, 0.32, "≤ window", size=11, color=A5, align=C)
text(s, 6.75, 4.45, 5.9, 0.3, "Log buffer: one memory-mapped file, three rotating terms", size=11, color=A5)
card(s, 6.75, 5.0, 2.85, 1.2, [{"text": "Tethered", "bold": True, "size": 15},
                                {"text": "The publication waits for it.", "size": 13, "color": A5}])
card(s, 9.85, 5.0, 2.88, 1.2, [{"text": "Untethered", "bold": True, "size": 15},
                                {"text": "Dropped when a window behind; may rejoin.", "size": 13, "color": A5}])

s = content("The media driver", "Applications do no networking: a driver owns sockets and log buffers",
            "Embedded in a process or standalone (aeronmd). Clients reach it through the command-and-control file in a "
            "shared Aeron directory. Conductor: client commands, log buffer lifecycle, counters, timeouts. Sender: "
            "reads publication terms, sends data frames, retransmits on NAK, respects the receivers' window. Receiver: "
            "reassembles frames into an image per remote publisher, NAKs gaps, sends status messages that advertise "
            "the window. Over IPC there is no sender or receiver: the subscriber reads the publisher's log buffer "
            "directly. Each seqeron member embeds its driver; clients on the host attach to it.")
f = Fig(s, L, 1.8, CW / 960)
f.host(10, 60, 160, 200, "Publisher process")
f.box(25, 110, 130, 50, "Publication", "offer() · tryClaim()")
f.host(195, 40, 260, 240, "Media driver")
f.box(210, 78, 230, 56, "Publication log buffer", "mmap file · 3 terms", "mem")
f.box(210, 155, 110, 52, "Conductor", "commands · counters")
f.box(330, 155, 110, 52, "Sender", "send · retransmit")
f.label(195, 228, 260, 30, "clients attach via cnc.dat")
f.host(505, 40, 260, 240, "Media driver")
f.box(520, 78, 230, 56, "Image log buffer", "one per remote publisher", "mem")
f.box(520, 155, 110, 52, "Receiver", "reassemble · NAK")
f.box(640, 155, 110, 52, "Conductor", "commands · counters")
f.label(505, 228, 260, 30, "clients attach via cnc.dat")
f.host(790, 60, 160, 200, "Subscriber process")
f.box(805, 110, 130, 50, "Subscription", "poll() fragments")
f.arrow([(155, 128), (210, 108)])
f.arrow([(385, 134), (385, 155)])
f.arrow([(440, 170), (520, 170)], "a")
f.label(445, 146, 70, 22, "data", color=A1)
f.arrow([(520, 192), (440, 192)], "d")
f.label(440, 196, 80, 22, "SM · NAK")
f.arrow([(575, 155), (575, 134)])
f.arrow([(750, 108), (805, 128)])
text(s, L, 6.05, CW, 0.6, [{"runs": [("Over aeron:ipc ", {"bold": True, "color": TX}),
                                     ("publisher and subscriber share one log buffer: no sender, no receiver, no socket. "
                                      "Every stage is visible as a position counter.", {})]}], size=14, color=A5)

s = content("Aeron Archive", "Streams recorded to local files, replayed from any position",
            "The archive records any stream to segment files and keeps a catalog of recordings with ids and start and "
            "stop positions. A replay starts from any position onto any channel and stream as a new session, and a "
            "replay of a recording that is still being written follows it live. The RecordingPos counter shows how far "
            "a recording has got: seqeron's queue publisher watches exactly that counter to decide whether its archive is "
            "slow, stalled or gone.")
segs = 5
for i in range(segs):
    x = L + i * (CW / segs)
    box(s, x, 2.15, CW / segs, 0.75, None, None, style="queue")
    text(s, x, 2.15, CW / segs, 0.75, f"segment file {i}", size=12, color=A1, align=C, anchor=M)
text(s, L, 1.7, 3, 0.35, "start position", size=12, color=A5)
text(s, W - L - 3, 1.7, 3, 0.35, "RecordingPos (live) →", size=12, color=A1, align=R, bold=True)
arrow(s, [(5.5, 2.9), (5.5, 3.45), (W - L, 3.45)], "w")
text(s, 5.65, 3.5, 7.0, 0.35, "replay(recordingId, position, length): a new session on any channel and stream",
     size=12, color=A2)
for i, (t, d) in enumerate([("Catalog", "Every recording has an id and start and stop positions. Nothing indexes by message."),
                            ("Replay", "From any position, as a new session. A bounded replay of an active recording follows it live."),
                            ("Control", "A request/response protocol drives it; counters expose each recording's progress.")]):
    x = L + i * (3.84 + 0.305)
    card(s, x, 4.3, 3.84, 2.25, [{"text": t, "bold": True, "size": 18, "head": True, "space": 6},
                                 {"text": d, "size": 14, "color": TX}])

s = content("Sequencer components", "The replicated state machine has no Aeron dependency",
            "Sequencer owns globalSeqNo, the gateway list and session bindings, promotion deadlines and the snapshot "
            "policy; it validates and encodes every frame and hands each to a sink in globalSeqNo order, the event's own "
            "frame first and whatever it triggered behind it. SequencerService is the ClusteredService adapter: the sink "
            "is the queue, and it holds no replicated state and decides no frame order. QueuePublisher is the reliable-offer discipline for "
            "the queue. The rule underneath all of it is spec S-3: every reject decision and every byte of every "
            "synthesized frame is a function of the committed log alone.")
parts = [("SequencerService", "ClusteredService adapter", "Decides when to call the sequencer; its queue is the frame sink. Holds no replicated state."),
         ("Sequencer", "the replicated state machine", "Owns globalSeqNo, gateway list, bindings, deadlines. Validates, encodes and publishes in order."),
         ("QueuePublisher", "reliable, bounded offer", "Retries the queue offer; watches RecordingPos; terminates the member on a stalled archive.")]
for i, (t, k, d) in enumerate(parts):
    x = L + i * (3.84 + 0.305)
    card(s, x, 1.85, 3.84, 2.2, [{"text": t, "bold": True, "size": 18, "head": True},
                                  {"text": k, "size": 13, "color": A1, "space": 8},
                                  {"text": d, "size": 14}], style="plain")
    if i < 2:
        arrow(s, [(x + 3.84, 2.95), (x + 3.84 + 0.305, 2.95)], "a")
card(s, L, 4.4, CW, 1.7, [{"text": "S-3  Determinism", "bold": True, "size": 13, "color": BG, "space": 6},
                           {"text": "Every reject decision and every byte of every synthesized frame is a function of "
                                    "the committed log alone: not node configuration, wall-clock time, local state or "
                                    "hash-map iteration order.", "size": 17, "color": BG, "head": True}],
     style="dark", margin=0.3, anchor=M)

s = content("Frames created by the sequencer", "Four synthesized frames share the one globalSeqNo counter",
            "These carry sourceId, connectionId and sessionId of −1 and the consensus timestamp. ClusterHeartbeat is "
            "driven by a cluster timer, so it is the same on every member and advances while every producer is silent, "
            "which is exactly when a watchdog needs a clock. LeadershipChanged is emitted per term even when the same "
            "member wins again. GatewayActive is the election's decision. SnapshotStarted is the cut of an application "
            "snapshot round.")
frames = [("ClusterHeartbeat", "104", "Every second, from a cluster timer.", "A clock every client agrees on; keeps advancing while producers are silent."),
          ("LeadershipChanged", "105", "Once per Raft term, same leader or not.", "Gates leader-only work; tells producers which ingress was lost."),
          ("GatewayActive", "106", "Bootstrap, session close, timeout, operator.", "The gateway election's decision, at one position for everyone."),
          ("SnapshotStarted", "107", "Operator request or policy interval.", "The cut of an application snapshot round.")]
for i, (t, tid, when, why) in enumerate(frames):
    col, row = i % 2, i // 2
    x, y = L + col * (5.9 + 0.33), 1.85 + row * (2.3 + 0.25)
    shape(s, x, y, 5.9, 2.3, "syn", name="Frame card")
    text(s, x + 0.25, y + 0.18, 4.4, 0.5, t, size=20, bold=True, head=True)
    text(s, x + 4.65, y + 0.2, 1.0, 0.45, tid, size=14, color=A2, align=R, bold=True)
    text(s, x + 0.25, y + 0.8, 5.4, 0.6, [{"runs": [("When  ", {"bold": True}), (when, {})]}], size=14)
    text(s, x + 0.25, y + 1.35, 5.4, 0.8, [{"runs": [("Why  ", {"bold": True}), (why, {})]}], size=14)

s = content("The replay protocol", "Three node-local messages; never sequenced, never recorded",
            "Requests on stream 202, replies on 203, replay data on 201. A client addresses itself by a clientId it "
            "draws once from its media driver, so no two clients share one. Every request carries a requestId the client "
            "increments on every send, including an unchanged resend, and the replayer echoes; a client ignores any reply "
            "that is not for its current request. The client detects completion by reaching catchUpPosition, not by the "
            "replay image closing, because a bounded replay of an active recording does not close at the bound. There is "
            "no heartbeat and no completion message: the Replayer frees a slot when the replay's publication leaves the "
            "driver's counters, which the client's closing its replay subscription does at once, and the driver's "
            "liveness timeout does for a client that died.")
table(s, L, 1.85, [3.9, 2.2, 6.03], [
    ["Message", "Direction", "Meaning"],
    ["ReplayRequest", "client → replayer", "From the start, or resume at a recording position"],
    ["Replaying", "replayer → client", "Follow this replay session to catchUpPosition"],
    ["ReplayPending", "replayer → client", "Not now: no free slot, or the position is not recorded yet"],
], size=14, rowh=0.55)
card(s, L, 4.3, CW, 1.15, [{"runs": [("A slot lives as long as its replay's publication. ", {"bold": True}),
                                      ("The client frees it by closing its replay subscription at catchUpPosition; a "
                                       "client that dies frees it within the driver's 10 s liveness timeout.", {})]}],
     anchor=M)
text(s, L, 5.75, CW, 0.4, "clientId comes from the client's media driver; requestId is echoed on every reply, and a client "
     "ignores replies to anything but its current request.", size=13, color=A5)

LIVE_LANES = [("Live queue", "stream 205", "queue"), ("Client", "ReplayerStreamReceiver", "box"),
              ("ReplayerService", "requests 202 · replies 203", "syn"), ("Archive", "active queue recording", "plain")]

SEQ_X = [1.7, 4.1, 8.4, 11.65]

s = content("Replay sequence: a cold start", "Walk the recording to its end, then follow the queue",
            "The Replayer is ready once its active recording is visible. A client's cold start asks for the active "
            "recording FROM_START, under the clientId it drew from its media driver; the Replayer starts a replay "
            "bounded at the recording's end as it is now and answers Replaying with the replay session and "
            "catchUpPosition, echoing the requestId. The client subscribes to stream 201 for that session only, so it "
            "never holds up another client's replay. Its first frame must be globalSeqNo 1, or the client stops: the "
            "recording is short. Live queue frames past the end arrive meanwhile and go into the retained FIFO. Completion "
            "is by position, not by the image closing: a bounded replay of a recording still being written does not "
            "close at its bound. Closing the replay subscription frees the slot; the FIFO drains, de-duplicated on "
            "globalSeqNo, and the client is caught up. An empty recording answers NO_REPLAY_NEEDED and the client "
            "follows the queue at once.")
q = Seq(s, LIVE_LANES, SEQ_X)
q.msg(1, 2, "ReplayRequest(clientId, requestId 1, FROM_START)")
q.msg(2, 3, "startReplay(start → last recorded) · 201")
q.msg(2, 1, "Replaying(session, catchUpPosition = last recorded)", "d")
q.note(1, "subscribe to 201 for this replay session only")
q.msg(3, 1, "frame 1: first frame must be globalSeqNo 1", "w")
q.begin("LOOP", "until the replay image reaches catchUpPosition")
q.msg(3, 1, "replayed frames 2, 3 … → dispatched in globalSeqNo order", "w")
q.msg(0, 1, "frames past the end → retained", "a")
q.end()
q.note(1, "position ≥ catchUpPosition: close the replay")
q.note(2, "publication gone → slot freed for the next request")
q.note(1, "drain the FIFO → isCaughtUp(), onCaughtUp", style="queue")

s = content("Replay sequence: a gap", "Resume at the last delivered frame; the anchor checks the recording",
            "A slow consumer is dropped by its driver and misses 105. Frame 106 arrives ahead of the next expected number: "
            "isCaughtUp() clears, 106 is retained, and the client asks to resume at the position where 104 ended, the "
            "last frame it delivered. The resumed replay's first frame must be 105, the anchor. If it is any other frame, "
            "or the Replayer refuses the position with NO_REPLAY_NEEDED, the client stops: positions follow from the "
            "frames, and the recording changes only with the member, whose restart takes the client's driver with it. "
            "A position not recorded yet is answered ReplayPending. Live frames keep arriving into the FIFO, bounded at "
            "65,536 frames or 16 MiB; past either bound the client drops them and resumes again at its last delivered "
            "frame when the replay ends. On reaching catchUpPosition the client closes the replay, the FIFO drains 106 "
            "onward and the client is caught up, with no second gap.")
q = Seq(s, LIVE_LANES, SEQ_X)
q.msg(0, 1, "104 → dispatched", "a")
q.msg(0, 1, "105", lost=True)
q.msg(0, 1, "106, ahead of 105", "a")
q.note(1, "isCaughtUp() cleared; 106 retained")
q.msg(1, 2, "ReplayRequest(requestId k, fromPosition = end of 104)")
q.msg(2, 3, "startReplay(fromPosition → last recorded)")
q.msg(2, 1, "Replaying(session, catchUpPosition = last recorded)", "d")
q.msg(0, 1, "107, 108 … → retained", "a")
q.begin("ALT", "the first replayed frame is 105, the anchor")
q.msg(3, 1, "105 → dispatched", "w")
q.note(1, "catchUpPosition: close the replay, slot freed")
q.note(1, "drain 106, 107, 108 → isCaughtUp(), onCaughtUp", style="queue")
q.orelse("another frame, or NO_REPLAY_NEEDED: the recording does not match")
q.note(1, "fatal: raised from the client's duty cycle", style="syn")
q.end()

s = content("Promotion rules", "Four triggers, one target rule, all on consensus time",
            "Target rule: the lowest preference rank among siblings, excluding the instance that lost the role, so a "
            "pair whose instances were all down converges on whichever starts first. The activation timeout arms at the "
            "GatewayActive's timestamp plus 5 s on the consensus clock and is evaluated on each heartbeat, which hands "
            "over every overdue deadline, each gateway at most once, so repeated failures move down the list. Being superseded is not a failure: the "
            "instance closes its edge, keeps its session and keeps following the queue so it can be designated again. From "
            "the GatewayActive that superseded it the sequencer sequences nothing from its session (rule S-6), so it drops "
            "what it still had unconfirmed, and its session closing later promotes nothing. A fence is different: the "
            "instance no longer trusts its own view: its façade closes the edge through onStandby, releases its session, "
            "reports once and does nothing more, which promotes the standby at the fence. The application decides "
            "when to exit; the repository's convention is code 70.")
trig = [("Bootstrap", "After the last topology row, each rank-0 instance."),
        ("Session close", "The designated instance's bound session closes."),
        ("Activation timeout", "No GatewayStarted within 5 s of consensus time."),
        ("Operator", "clusterctl activate names an instance directly.")]
tw = (CW - 3 * 0.3) / 4
for i, (t, d) in enumerate(trig):
    x = L + i * (tw + 0.3)
    shape(s, x, 1.85, tw, 2.3, "card", name="Trigger")
    num_circle(s, x + 0.2, 2.05, i + 1, d=0.45, size=14)
    text(s, x + 0.2, 2.65, tw - 0.4, 0.45, t, size=17, bold=True, head=True)
    text(s, x + 0.2, 3.1, tw - 0.4, 1.0, d, size=14)
text(s, L, 4.35, CW, 0.4, [{"runs": [("Target: ", {"bold": True}),
                                     ("lowest rank among siblings, excluding the instance that lost the role.", {})]}], size=15)
card(s, L, 4.95, 5.9, 1.7, [{"text": "Superseded → standby", "bold": True, "size": 17, "head": True, "space": 6},
                            {"text": "Closes its edge, keeps its session, keeps following the queue. Nothing it sends is sequenced. Can be designated again.", "size": 14}],
     style="plain")
card(s, L + 6.23, 4.95, 5.9, 1.7, [{"text": "Fenced → session released", "bold": True, "size": 17, "head": True, "color": A2, "space": 6},
                                   {"text": "Its view of the log can no longer be trusted. Closes its edge, releases its session, reports once; the standby is promoted at the fence.", "size": 14}],
     style="syn")

s = content("Applications: cluster member local or remote", "Each application follows its own member's queue; only the leader's publishes",
            "An application runs one instance on every member, attached to that member's media driver, and a "
            "member can host several applications, each with its own sourceId and client id. Every instance follows "
            "its own member's queue and so holds the same state. LeadershipChanged picks the instance that acts: only the one on "
            "the leader's member publishes, over IPC ingress to its own consensus module, and every leadership change "
            "closes its gate, so a reply lost in the election is dispatched again. Remote, the application is the one "
            "instance on its own host, with no member there. Like a remote gateway's host, that host runs a local "
            "ReplayerService whose QueueRelay copies a member's queue onto the host's own queue and archive. The application's gate "
            "opens once caught up, whoever leads, and it submits over UDP. Nothing elects between two remote instances, "
            "so a remote service that needs a standby is a gateway.")
for i in range(3):
    y = 1.95 + i * 1.2
    shape(s, L, y, 3.3, 1.05, "host", name="Member")
    text(s, L + 0.12, y + 0.04, 3.0, 0.3, f"MEMBER {i}" + (" · LEADER" if i == 0 else ""), size=10, bold=True,
         color=A5)
    box(s, L + 0.15, y + 0.35, 1.25, 0.55, "Application", "publishes" if i == 0 else "follows",
        style="lead" if i == 0 else "box", tsize=12, ssize=10)
    box(s, L + 1.9, y + 0.35, 1.25, 0.55, "Sequencer", style="box", tsize=12)
    arrow(s, [(L + 1.9, y + 0.55), (L + 1.4, y + 0.55)], "a")
    if i == 0:
        arrow(s, [(L + 1.4, y + 0.75), (L + 1.9, y + 0.75)])
shape(s, 4.8, 1.95, 3.3, 2.25, "host", name="Application host")
text(s, 4.92, 1.99, 3.0, 0.3, "REMOTE APPLICATION HOST", size=10, bold=True, color=A5)
box(s, 4.95, 2.3, 3.0, 0.65, "Application", "one instance · gate open once caught up", style="box", tsize=13, ssize=10)
box(s, 4.95, 3.35, 3.0, 0.65, "ReplayerService", "local · QueueRelay: the host's own queue", style="syn", tsize=13,
    ssize=10)
arrow(s, [(6.45, 3.35), (6.45, 2.95)], "a")
arrow(s, [(4.95, 2.62), (L + 3.3, 2.62)])
text(s, L + 3.3, 2.33, 0.9, 0.25, "UDP ingress", size=10, color=A5, align=PP_ALIGN.CENTER)
arrow(s, [(L + 3.3, 3.67), (4.95, 3.67)], "a")
text(s, L + 3.3, 3.38, 0.9, 0.25, "replay", size=10, color=A1, align=PP_ALIGN.CENTER)
text(s, L, 5.75, 7.5, 0.8, [
    "Teal: the queue, live.",
    "Black: ingress, IPC from the application on the leader and UDP from the remote host."], size=12, color=A5, space=2)
text(s, 8.6, 1.95, 4.13, 4.5, [
    {"text": "An application on every member; a member hosts several", "bullet": True, "space": 10},
    {"text": "Every application follows its own member's queue and holds the same state", "bullet": True, "space": 10},
    {"text": "Only the application on the leader publishes; every LeadershipChanged closes its gate", "bullet": True, "space": 10},
    {"text": "Remote: one instance on its own host, beside a local ReplayerService; UDP ingress", "bullet": True},
], size=15)

s = content("Remote hosts: the queue relay", "Clients off the cluster see exactly what they would see on a member",
            "A gateway or application host, with clients but no member, runs its own ReplayerServer with a media driver and archive. One archive "
            "replay over UDP follows a member's active recording live and republishes it on the host's own queue. QueueRelay "
            "republishes only the frame after the last one it published, so the host's queue is the log frame for frame "
            "whichever member it came from. After 3 s with no frame (three heartbeats), or a failed read, it fails over "
            "to the next member archive in replayer.archiveEndpoints, round robin; any member serves, leader or not. It "
            "resumes at the same recording position, valid because every member's active recording starts at globalSeqNo 1 and positions "
            "follow from the frames; if the frame there is not the next one it replays that recording from its start "
            "and skips what it has. A restarted relay relays from 1 into a new local recording.")
for i in range(3):
    y = 1.95 + i * 1.3
    box(s, L, y, 2.6, 0.95, f"Member {i}", "active recording", style="box" if i else "lead", tsize=14, ssize=11)
arrow(s, [(L + 2.6, 2.42), (4.15, 3.55)], "a")
arrow(s, [(L + 2.6, 3.72), (4.15, 3.75)], "d")
text(s, L, 5.75, 7.5, 0.8, [
    "Solid: one UDP archive replay.",
    "Dashed: failover to next member archive."], size=12, color=A5, space=2)
shape(s, 4.2, 1.95, 3.9, 3.55, "host", name="Remote host")
text(s, 4.35, 2.02, 3.6, 0.35, "GATEWAY OR APPLICATION HOST", size=11, bold=True, color=A5)
box(s, 4.4, 2.5, 3.5, 0.75, "QueueRelay", "next frame only", style="syn", tsize=14, ssize=11)
text(s, 4.95, 3.33, 1.9, 0.24, "Queue · aeron:ipc 205", size=10, bold=True, color=A1, anchor=MSO_ANCHOR.MIDDLE)
log_cells(s, 4.95, 3.6)
box(s, 4.4, 4.5, 1.4, 0.75, "Archive", "records the queue", style="box", tsize=14, ssize=11)
box(s, 6.0, 4.5, 1.9, 0.75, "Clients", "unchanged; UDP ingress", style="box", tsize=14, ssize=11)
arrow(s, [(6.489, 3.25), (6.489, 3.6)])
arrow(s, [(5.451, 3.96), (5.451, 4.5)])
arrow(s, [(6.143, 3.96), (6.143, 4.5)], "a")
text(s, 8.6, 1.95, 4.13, 4.5, [
    {"text": "Republishes only globalSeqNo last + 1", "bullet": True, "space": 10},
    {"text": "Resumes at the same recording position on the next member", "bullet": True, "space": 10},
    {"text": "Restart relays from globalSeqNo 1 into a new recording, as a member's restart does", "bullet": True, "space": 10},
    {"text": "Costs the member one archive replay per host; the sequencer nothing", "bullet": True},
], size=15)

s = content("A lost session is replaced, not fenced", "Unless the instance announced itself on it",
            "The cluster closes a client's session when its keep-alives stop: a paused process, a lost quorum. Closing "
            "the session a designated instance's GatewayStarted bound is what promotes its sibling, and the queue carries no "
            "frame for that close, so a new session could act before the instance knew whether it had been superseded. "
            "That instance fences. Every other client, an application, a standby, or an instance designated after "
            "its session was lost, opens a new session the way it opened the first: an attempt at most once a second, "
            "advanced a step per doWork(), so the duty cycle never waits on an unreachable cluster. It fences only if no "
            "session opens within 20 s and the attempt then under way fails. What was unconfirmed on the old session is "
            "dropped, not resent. session-replace-test.sh drives it through a paused client, a lost quorum and a paused "
            "gateway pair, in Java.")
for i, (t, d, st, col) in enumerate([
        ("Replaces", "An application, a gateway standby, or an instance designated after its session was lost. "
                     "It opens a new session the way it opened the first.", "card", A1),
        ("Fences", "A gateway instance that announced itself on the lost session. Closing that session promotes its "
                   "sibling, and the queue carries no frame for the close.", "syn", A2),
        ("Drops", "What was unconfirmed on the old session is not resent: its close is on no queue, so what reached the "
                  "log cannot be counted.", "card", TX)]):
    x = L + i * (3.84 + 0.305)
    card(s, x, 1.85, 3.84, 3.1, [{"text": t, "size": 20, "bold": True, "head": True, "color": col, "space": 8},
                                 {"text": d, "size": 15}], style=st, margin=0.25)
card(s, L, 5.3, CW, 1.3, [{"runs": [("Never blocks the duty cycle. ", {"bold": True}),
                                    ("An attempt starts at most once a second and advances a step per doWork(). With no "
                                     "session after 20 s and the attempt under way failed, the instance fences.", {})],
                           "size": 16}], margin=0.25, anchor=M)

s = content("The snapshot protocol", "Four frames on the log; the records never enter it",
            "The topology turns rounds on: SnapshotPolicyRegistered sets the interval, 0 meaning operator requests only, "
            "and each source's row says whether it takes part, which its façades read and the sequencer does not. The "
            "sequencer decodes the policy alone and answers a SnapshotRequested with SnapshotStarted at the next "
            "globalSeqNo, as a GatewayActive follows a GatewayActivationRequested. SnapshotEnd belongs to the frame's "
            "sourceId, one per application or logical gateway per round; the sequencer checks only its format, since "
            "rejecting a late or superseded end would fence its producer. Each instance confirms its own file against "
            "the sequenced end, and the file's trailer records the round's cut, which a restore checks against the "
            "recording. The Replayer holds no snapshot state, and the records themselves never enter the log.")
table(s, L, 1.85, [2.85, 2.6, 3.0, 3.68], [
    ["Message", "Sent by", "Fields", "Effect"],
    ["SnapshotPolicyRegistered", "clusterctl load-topology", "intervalSeconds", "Turns rounds on; 0 is operator requests only"],
    ["GatewayRegistered, ApplicationRegistered", "clusterctl load-topology", "snapshot", "Whether that source takes part; read by its façades"],
    ["SnapshotRequested", "clusterctl request-snapshot", "correlationId", "SnapshotStarted at the next globalSeqNo"],
    ["SnapshotStarted", "the sequencer only", "round", "The cut: every participating instance serializes here"],
    ["SnapshotEnd", "the instance that may publish", "round, recordCount, length, crc32c", "Each instance compares: a match confirms its file, a difference fences"],
], size=13, rowh=0.55)
text(s, L, 5.55, CW, 0.4, "All schema 210. A file's trailer carries its round, record count, length, CRC32C and cut: "
     "globalSeqNo, recording position and timestamp.", size=13, color=A5)

s = content("Replay sequence: when history is not served", "Pending holds the client at the gap; a mismatch stops it",
            "Every reply echoes the requestId; a client ignores any reply that is not for its current request. "
            "ReplayPending means the request is pending or the Replayer cannot start it yet: four replays at most, one "
            "while the queue's recording trails the queue by more than a quarter of the publication window, not ready "
            "before its active recording is visible, a resume position the recording has not reached yet, or an "
            "archive call that threw, which is counted. The client's resend every 500 ms is the retry, and a pending "
            "request starts once a slot frees. A resume the Replayer refuses, NO_REPLAY_NEEDED for a position, means "
            "the recording does not match the positions the client holds, and the client stops, as it does when a "
            "resume opens on the wrong frame. A replay that stops under the client, no progress for 5 s or an image "
            "closed short of catchUpPosition, is requested again; a resume re-anchors. Nothing dispatched for 30 s "
            "raises the recoveryStalled gauge; a façade fences RECOVERY_STALLED after 60 s on a designated instance or "
            "the leader's replica, above the driver's 10 s liveness timeout that frees a dead client's slot.")
q = Seq(s, LIVE_LANES, SEQ_X)
q.msg(1, 2, "ReplayRequest(requestId n, fromPosition)")
q.begin("ALT", "no free slot (4, or 1 while the recorder trails), not ready, not recorded yet, or an archive error")
q.msg(2, 1, "ReplayPending(n): hold at the gap", "d")
q.msg(1, 2, "ReplayRequest(n + 1, same position) after 500 ms", )
q.orelse("the position is not a frame in the recording")
q.msg(2, 1, "Replaying(n, NO_REPLAY_NEEDED)", "d")
q.note(1, "fatal: raised from the client's duty cycle", style="syn")
q.end()
q.begin("ALT", "an established replay stops: no progress for 5 s, or its image closes short")
q.msg(1, 2, "ReplayRequest(new requestId); a resume re-anchors")
q.end()
q.note(1, "30 s with nothing dispatched: recoveryStalled gauge")
q.note(1, "60 s: the façade fences RECOVERY_STALLED", style="syn")

s = content("Replay sequence: a snapshot restore", "Check the file's cut on the recording, restore, then resume after it",
            "A client whose façade has a SnapshotListener restores before it dispatches anything. It opens its confirmed "
            "file, whose trailer records the round's cut: globalSeqNo, recording position and timestamp. It replays "
            "from that position, and the first frame must be the round's SnapshotStarted at that globalSeqNo and "
            "timestamp. Any other frame, or NO_REPLAY_NEEDED, means the file was cut from another log, after a purge "
            "for example, and it gives way to an older file; with none left the client walks from globalSeqNo 1. "
            "ReplayPending holds it while a restarted member re-records. The records are read from the local file, at "
            "most 1,024 per poll. A format or header version the build does not read, or records that fail the "
            "trailer, halt recovery: restoreFailure, and the façade fences SNAPSHOT_UNRESTORABLE. The resume after it "
            "is an ordinary one, anchored on the cut, which is dropped as already delivered.")
q = Seq(s, [("Snapshot files", "SnapshotStore, local", "card")] + LIVE_LANES[1:3]
        + [("Archive", "active queue recording", "plain")], SEQ_X)
q.msg(0, 1, "newest round r and its cut")
q.msg(1, 2, "ReplayRequest(cut position), anchor = the cut")
q.msg(2, 1, "Replaying(session, catchUpPosition)", "d")
q.msg(3, 1, "first frame: SnapshotStarted(r) at the cut?", "w")
q.begin("ALT", "another frame, or NO_REPLAY_NEEDED: a file from another log")
q.msg(1, 0, "the next older round")
q.note(1, "none left: ReplayRequest(FROM_START), a walk")
q.end()
q.msg(0, 1, "records, ≤ 1,024 per poll")
q.note(1, "lastGlobalSeqNo = cut; leadership from the header")
q.msg(1, 2, "ReplayRequest(cut position), anchor = the cut")
q.msg(2, 1, "Replaying(session, catchUpPosition = last recorded)", "d")
q.msg(3, 1, "SnapshotStarted (the cut, dropped), then the frames after it", "w")
q.note(1, "close the replay, drain the FIFO → caught up", style="queue")

s = content("Snapshot sequence: a round", "Cut, serialize, end, compare; the records never enter the log",
            "A round starts on an operator's SnapshotRequested or when the policy interval elapses at a heartbeat; "
            "either way the sequencer synthesizes SnapshotStarted, and its globalSeqNo is the cut. Every instance of a "
            "participating source serializes its state there into its own directory, standby and follower replicas "
            "included. Only the instance whose gate is open, the designated gateway instance or the leader's replica, "
            "submits SnapshotEnd, and the sequencer checks only its format. Every instance then compares its own "
            "serialization with the sequenced end; one that differs no longer holds the state the log implies and "
            "fences; a match confirms its file, which is forced to disk, renamed <round>.snapshot and replaces every "
            "older one. A new round supersedes this one: an end not yet placed is abandoned.")
q = Seq(s, [("clusterctl", "request-snapshot", "plain"), ("Sequencer", "every member", "lead"),
            ("Instance A", "gate open: publishes", "box"), ("Instance B", "every other instance", "box")],
        [1.9, 5.1, 8.3, 11.5], legend=(("queue", "a"), ("ingress", "ln")))
q.msg(0, 1, "SnapshotRequested")
q.note(1, "or: interval, at a heartbeat", w=2.6)
q.msg(1, 3, "SnapshotStarted(r) on every queue: the cut R", "a")
q.note(2, "each serializes at R to its own <r>.tmp", w=3.8)
q.msg(2, 1, "SnapshotEnd(r, count, length, crc)")
q.msg(1, 3, "SnapshotEnd sequenced, on every queue", "a")
q.note(2, "each compares its own file with the end", w=3.8)
q.begin("ALT", "an instance's file matches the end")
q.note(2, "it confirms: <r>.snapshot, older files deleted", w=3.8, style="queue")
q.orelse("it differs")
q.note(2, "it fences: SNAPSHOT_DIVERGED", w=3.8, style="syn")
q.end()

s = content("The gateway protocol","Six system frames; only the sequencer designates",
            "This is seqeron's vocabulary, not any edge protocol's. clusterctl load-topology puts the list on the log, one "
            "GatewayRegistered per instance; an instance finds its own row by name and publishes nothing before it has. "
            "GatewayActive is never submitted: the sequencer synthesizes it, and an operator's activate is a request it "
            "validates first. The designated instance answers with GatewayStarted, which binds its cluster session and "
            "carries the first connection id it will allocate, above any it saw, so ids never repeat across a restart. "
            "Connection frames are the generic half: ConnectionOpened carries the producer's own connection data, opaque "
            "to the cluster, so a standby can map an id to its session before that connection's first message. The "
            "sequencer keeps no connection state: a consumer treats GatewayStarted as closing the pair's open "
            "connections. S-6 is "
            "condition 10 of validation, and every input to it comes from the log.")
table(s, L, 1.85, [2.95, 2.5, 3.0, 3.68], [
    ["Frame", "Sent by", "Fields", "The sequencer"],
    ["GatewayRegistered", "clusterctl load-topology", "gatewayId, sourceId, name, rank, remaining", "Builds the list; remaining 0 bootstraps rank 0"],
    ["GatewayActive", "the sequencer only", "gatewayId", "Designates one instance; arms a 5 s deadline"],
    ["GatewayStarted", "the designated instance", "gatewayId, firstConnectionId", "Binds its session"],
    ["GatewayActivationRequested", "clusterctl activate", "gatewayId", "Validates it, then a GatewayActive at the next number"],
    ["ConnectionOpened", "the active instance", "connectionData, opaque", "Forwards it unopened"],
    ["ConnectionClosed", "the active instance", "none", "Forwards it"],
], size=13, rowh=0.52)
card(s, L, 5.75, CW, 0.95, [{"runs": [("The log is the judge (S-6). ", {"bold": True}),
                                      ("A GatewayStarted must name the instance the latest GatewayActive designated. Once a "
                                       "GatewayActive names a sibling, nothing from the old instance's session is sequenced, "
                                       "even while that session stays open.", {})], "size": 14}],
     margin=0.22, anchor=M)

s = content("The application protocol", "What one replica reads and submits",
            "An application brings its own protocol in the application family. payloadId names one schema or encoding, "
            "allocated by the deployment from 2 up; 0 is invalid and 1 is retired. The producer encodes the payload once "
            "and nothing re-encodes it, so any encoding works and its versioning is the application's own. The table is "
            "what Application does on the application's behalf, in every language. It submits application frames only "
            "while its leader gate is open; a reply goes out under the requester's sourceId and connectionId so the "
            "gateway that took the request can route it. Every frame on the queue also passes confirmed ingress, which "
            "matches the replica's own frames by cluster session. LeadershipChanged reaches the gate and confirmed "
            "ingress, never the application directly; the listener sees only the gate's edges. The consumer rules keep "
            "applications independent: a consumer skips what it does not recognise, and since continuity rests on "
            "globalSeqNo, skipping can neither create nor hide a gap. A PayloadIdRegistered row in the topology is a "
            "label for SbeLogPrinter, not permission; a per-node allowlist would let nodes diverge. The snapshot frames "
            "it reads and submits are on the snapshot protocol slide.")
table(s, L, 1.85, [2.95, 1.2, 7.98], [
    ["Frame", "Direction", "The façade"],
    ["Sequenced", "reads", "onSequenced, in sequence number order; its own frames confirm its sends"],
    ["Unsequenced", "submits", "publish under its own sourceId; reply under the requester's. Only while the gate is open"],
    ["LeadershipChanged", "reads", "Closes the leader gate; a newer term resends what the old leader lost"],
    ["ClusterHeartbeat", "reads", "onClusterHeartbeat, on consensus time; resets the queue-stall fence"],
], size=14, rowh=0.45)
text(s, L, 4.22, CW, 0.4, "Snapshot frames are covered under the snapshot protocol. Every other system frame is "
     "skipped (P-2).", size=13, color=A5)
for i, (t, body) in enumerate([
        ("payloadId names a protocol", "One schema or encoding, never one message: SBE, protobuf, raw FIX. Encoded once "
                                       "by its producer and copied through unopened. Versioning is the application's."),
        ("Consumers skip what they do not know", "An unknown protocol or an unsupported version is skipped, not an "
                                                 "error. Continuity uses the sequence number, so a skip hides no gap. "
                                                 "Check the payload's own framing, an SBE schemaId, before decoding."),
        ("Registration labels; it does not admit", "PayloadIdRegistered names a protocol for SbeLogPrinter. The "
                                                   "sequencer rejects nothing for lacking a row. A protocol shared by "
                                                   "applications needs a row and one owning repository.")]):
    x = L + i * (3.84 + 0.305)
    card(s, x, 4.7, 3.84, 2.05, [{"text": t, "size": 15, "bold": True, "head": True, "space": 6},
                                 {"text": body, "size": 13}], margin=0.2)

s = content("Fences", "An instance that cannot trust its view stops before it acts on stale state",
            "Each fence is reported once through onFenced. Before it, the façade closes what lets the instance act, a "
            "serving gateway's edge through onStandby or an application's open gate through onLeadershipChanged(false), "
            "and releases its cluster session, so a designated instance's standby is promoted at the fence. At the end "
            "of that duty cycle it releases its queue and replay subscriptions, so a fenced instance holds no Replayer "
            "slot. Exiting is the application's decision; the examples exit with 70, the repository's convention for "
            "fenced, restart me. The queue-stall timer runs on local monotonic time because consensus time arrives in the very "
            "heartbeats being watched for. A replay is not a stall: the queue-stall check is suspended while not caught "
            "up, and the recovery-stall fence covers that time, armed only after the first catch-up because a cold start "
            "has no useful bound. 60 s exceeds the driver's 10 s liveness timeout, the longest a dead client holds a "
            "replay slot, so a correct Replayer cannot "
            "trip it. A stall fences only an instance that acts, the designated gateway instance or the leader's replica; "
            "any other logs it once, keeps following the queue and logs when it clears, since it cannot begin to act while "
            "stalled. A lost session fences only a gateway instance that announced itself on it; the slide on lost sessions "
            "covers the rest. Losing the media driver is not a fence; it raises from doWork().")
table(s, L, 1.85, [3.9, 8.23], [
    ["Fence", "Condition"],
    ["CLUSTER_SESSION_LOST", "Lost on a gateway instance that announced on it; elsewhere, not replaced in 20 s"],
    ["QUEUE_STALLED", "No ClusterHeartbeat for 20 s while caught up, on an instance that acts"],
    ["RECOVERY_STALLED", "Nothing recovered for 60 s after a first catch-up, on an instance that acts"],
    ["INGRESS_CONFIRM_FAULTED", "An own frame on the queue differs from the oldest pending one"],
    ["SNAPSHOT_DIVERGED", "Own snapshot differs from the one its source sequenced"],
    ["SNAPSHOT_UNRESTORABLE", "Newest confirmed snapshot cannot be restored"],
], size=14, rowh=0.58)
text(s, L, 6.1, CW, 0.5, "An instance that acts: the designated gateway instance, or the leader's replica. Any other logs a stall "
     "and keeps following.", size=13, color=A5)
text(s, L, 6.5, CW, 0.4, [{"runs": [("Before onFenced: ", {"bold": True}),
                                    ("the edge or leader gate closes and the session is released. Then exit, by "
                                     "convention with 70.", {})]}], size=13)

s = content("faulteron", "Network faults for Aeron, injected in eBPF",
            "faulteron is the injector behind the network-fault harnesses on the verification slide, in a repository of "
            "its own under Apache 2.0. A privileged container joins a member's network namespace and attaches two tc "
            "programs to its interface, one at egress and one at ingress, which share one rules map. The CLI writes "
            "rules into that map with bpftool. Each UDP packet goes to the most specific rule matching its peer, port, "
            "Aeron frame type, retransmission, or SBE schema and template; the rule drops a percentage, duplicates a "
            "percentage of the rest, and stamps a delay that the fq qdisc holds it to, so delay applies only on the "
            "way out. Process kills and pauses cannot produce these faults. Traffic over aeron:ipc never reaches an "
            "interface, so no rule touches it. Linux 5.18 or later, IPv4. seqeron's failover workflow runs its tests "
            "at a pinned commit.")
shape(s, L, 1.85, 7.15, 4.3, "frame")
text(s, L + 0.15, 1.9, 5.0, 0.3, "One container's network namespace", size=11, color=A5)
box(s, 0.85, 2.35, 1.7, 3.35, "seqeron member", "Aeron UDP: Raft, archive, ingress", tsize=14, ssize=11)
box(s, 2.95, 3.6, 1.8, 0.85, "faulteron CLI", "privileged; bpftool", style="syn", tsize=14, ssize=11)
box(s, 5.15, 2.35, 2.0, 0.85, "tc egress", "drop, duplicate, delay", style="queue", tsize=14, ssize=11)
box(s, 5.15, 3.6, 2.0, 0.85, "rules map", "peer, port, type, message", style="syn", tsize=14, ssize=11)
box(s, 5.15, 4.85, 2.0, 0.85, "tc ingress", "drop, duplicate", style="queue", tsize=14, ssize=11)
box(s, 7.5, 2.35, 0.5, 3.35, "eth0", style="host", tsize=11)
arrow(s, [(2.55, 2.775), (5.15, 2.775)])
arrow(s, [(7.15, 2.775), (7.5, 2.775)])
arrow(s, [(7.5, 5.275), (7.15, 5.275)])
arrow(s, [(5.15, 5.275), (2.55, 5.275)])
arrow(s, [(4.75, 4.025), (5.15, 4.025)], style="w")
arrow(s, [(6.15, 3.6), (6.15, 3.2)], style="d")
arrow(s, [(6.15, 4.45), (6.15, 4.85)], style="d")
arrow(s, [(8.0, 4.025), (8.4, 4.025)], both=True)
text(s, 2.95, 2.45, 1.8, 0.3, "out", size=11, color=A5, align=C)
text(s, 2.95, 4.95, 1.8, 0.3, "in", size=11, color=A5, align=C)
text(s, 8.45, 3.85, 0.75, 0.35, "peers", size=12, anchor=M)
card(s, 9.3, 1.85, W - L - 9.3, 4.3, [{"text": "What it can do", "size": 16, "bold": True, "head": True, "space": 8}]
     + [{"text": t, "size": 13, "bullet": True, "space": 6} for t in [
         "Cut one link, or a member off from its peers while its clients stay connected",
         "Take away one half of loss repair: NAKs or retransmissions",
         "Delay, jitter, reorder and duplicate a link",
         "Silence one member's votes and nothing else",
         "Not aeron:ipc, which never reaches an interface"]], margin=0.22)
url = text(s, L, 6.35, CW, 0.4, [{"runs": [("github.com/FredrikJDahlberg/faulteron", {"bold": True, "color": A1}),
                                           ("   Apache 2.0  ·  Linux 5.18 or later  ·  IPv4", {"color": A5})]}],
           size=14)
url.click_action.hyperlink.address = "https://github.com/FredrikJDahlberg/faulteron"

# ---------------------------------------------------------------- sections
if "sections" in OFF: SECTIONS.clear()
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
pres_el = prs.part._element
ext_lst = pres_el.find(qn("p:extLst"))
if ext_lst is None:
    ext_lst = etree.SubElement(pres_el, qn("p:extLst"))
ext = etree.SubElement(ext_lst, qn("p:ext"))
ext.set("uri", "{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
sl = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
for name, ids in SECTIONS:
    sec = etree.SubElement(sl, f"{{{P14}}}section")
    sec.set("name", name)
    sec.set("id", "{" + str(uuid.uuid4()).upper() + "}")
    lst = etree.SubElement(sec, f"{{{P14}}}sldIdLst")
    for sid in ids:
        etree.SubElement(lst, f"{{{P14}}}sldId").set("id", str(sid))

# python-pptx relates the notes master but never lists it; Keynote refuses such a package.
nm_rid = next(r.rId for r in prs.part.rels.values() if r.reltype == RT.NOTES_MASTER)
if pres_el.find(qn("p:notesMasterIdLst")) is None:
    nml = etree.Element(qn("p:notesMasterIdLst"))
    etree.SubElement(nml, qn("p:notesMasterId")).set(qn("r:id"), nm_rid)
    pres_el.find(qn("p:sldMasterIdLst")).addnext(nml)

prs.core_properties.title ="seqeron architecture briefing"
prs.core_properties.author = "Fredrik Dahlberg"
prs.save(OUT)
print("slides:", len(prs.slides))
